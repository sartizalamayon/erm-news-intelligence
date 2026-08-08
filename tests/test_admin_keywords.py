"""Keyword CRUD + rematch (app.services.admin, app.api.routers.keywords).
Service-layer tests against a real, migrated Postgres schema
(app.db.testing.migrated_schema); a couple of router-level tests for the
end-to-end HTTP wiring.
"""

from __future__ import annotations

import contextlib
import datetime as dt
import io

import pytest
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.api.deps import AuthenticatedUser, get_current_user
from app.db.queries import insert_article
from app.db.testing import migrated_schema, pg_available
from app.errors import ValidationError
from app.main import create_app
from app.services import admin as admin_service

pytestmark = pytest.mark.skipif(not pg_available(), reason="no reachable Postgres (set TEST_DATABASE_URL)")

ADMIN = AuthenticatedUser(id=9, email="admin@grameenphone.com", role="admin", display_name="Admin", session_id=9)


def test_create_list_get_keyword() -> None:
    with migrated_schema() as conn:
        created = admin_service.create_keyword(conn, section="telecom", term_en="spectrum", term_bn="স্পেকট্রাম")
        assert created["section"] == "telecom"
        assert created["enabled"] is True

        fetched = admin_service.get_keyword(conn, created["id"])
        assert fetched == created

        listed = admin_service.list_keywords(conn, section="telecom")
        assert [k["id"] for k in listed] == [created["id"]]


def test_create_keyword_requires_a_term() -> None:
    with migrated_schema() as conn:
        with pytest.raises(ValidationError):
            admin_service.create_keyword(conn, section="telecom", term_en="", term_bn="")


def test_create_keyword_english_only_term_is_valid() -> None:
    with migrated_schema() as conn:
        created = admin_service.create_keyword(conn, section="telecom", term_en="SMP", term_bn="")
        assert created["term_bn"] == ""


def test_list_keyword_sections_counts_per_section() -> None:
    with migrated_schema() as conn:
        admin_service.create_keyword(conn, section="telecom", term_en="spectrum")
        admin_service.create_keyword(conn, section="telecom", term_en="BTRC")
        admin_service.create_keyword(conn, section="macro", term_en="reserves")
        sections = {s["section"]: s["count"] for s in admin_service.list_keyword_sections(conn)}
        assert sections == {"telecom": 2, "macro": 1}


def test_update_keyword_partial_and_toggle() -> None:
    with migrated_schema() as conn:
        created = admin_service.create_keyword(conn, section="telecom", term_en="5G")
        updated = admin_service.update_keyword(conn, created["id"], term_bn="৫জি")
        assert updated["term_bn"] == "৫জি"
        assert updated["term_en"] == "5G"

        toggled = admin_service.toggle_keyword(conn, created["id"])
        assert toggled["enabled"] is False


def test_update_delete_unknown_keyword_returns_none_or_false() -> None:
    with migrated_schema() as conn:
        assert admin_service.update_keyword(conn, 999999, enabled=False) is None
        assert admin_service.delete_keyword(conn, 999999) is False


def test_delete_keyword_removes_row() -> None:
    with migrated_schema() as conn:
        created = admin_service.create_keyword(conn, section="telecom", term_en="NEIR")
        assert admin_service.delete_keyword(conn, created["id"]) is True
        assert admin_service.get_keyword(conn, created["id"]) is None


# ── KEYWORDS-CONFIG: custom sections, no-section keywords ──────────────


def test_create_keyword_with_a_brand_new_section_name_just_works() -> None:
    """`section` has never had a CHECK constraint or enum anywhere at
    the DB layer (migrations/0002_reference_tables.sql: plain TEXT) —
    this asserts that stays true: an admin can type ANY section name and
    it's simply stored, no migration or code change required per name.
    """
    with migrated_schema() as conn:
        created = admin_service.create_keyword(conn, section="Sports Diplomacy", term_en="cricket tour")
        assert created["section"] == "Sports Diplomacy"
        sections = {s["section"] for s in admin_service.list_keyword_sections(conn)}
        assert "Sports Diplomacy" in sections


def test_create_keyword_without_a_section_is_uncategorized() -> None:
    with migrated_schema() as conn:
        created = admin_service.create_keyword(conn, section=None, term_en="floating rate")
        assert created["section"] == "Uncategorized"
        created2 = admin_service.create_keyword(conn, section="", term_en="peg")
        assert created2["section"] == "Uncategorized"
        created3 = admin_service.create_keyword(conn, section="Uncategorized", term_en="repo rate")
        assert created3["section"] == "Uncategorized"

        listed = admin_service.list_keywords(conn, section="Uncategorized")
        assert {k["id"] for k in listed} == {created["id"], created2["id"], created3["id"]}

        sections = {s["section"]: s["count"] for s in admin_service.list_keyword_sections(conn)}
        assert sections == {"Uncategorized": 3}


def test_update_keyword_can_explicitly_clear_section_to_uncategorized() -> None:
    with migrated_schema() as conn:
        created = admin_service.create_keyword(conn, section="telecom", term_en="roaming")
        cleared = admin_service.update_keyword(conn, created["id"], section=None)
        assert cleared["section"] == "Uncategorized"


def test_rename_keyword_section_bulk_updates_matching_rows_only() -> None:
    with migrated_schema() as conn:
        a = admin_service.create_keyword(conn, section="telecom", term_en="spectrum")
        b = admin_service.create_keyword(conn, section="telecom", term_en="BTRC")
        other = admin_service.create_keyword(conn, section="macro", term_en="reserves")

        touched = admin_service.rename_keyword_section(conn, "telecom", "Telecom & Regulatory")
        assert touched == 2

        assert admin_service.get_keyword(conn, a["id"])["section"] == "Telecom & Regulatory"
        assert admin_service.get_keyword(conn, b["id"])["section"] == "Telecom & Regulatory"
        assert admin_service.get_keyword(conn, other["id"])["section"] == "macro"

        sections = {s["section"] for s in admin_service.list_keyword_sections(conn)}
        assert "telecom" not in sections
        assert "Telecom & Regulatory" in sections


def test_rematch_matches_keyword_with_no_section_assigned() -> None:
    """The matching pass (app.services.matching.load_keywords) must
    COALESCE a NULL section to a real string before grouping/sorting —
    otherwise `sorted({None, "telecom"})` raises TypeError the moment
    one uncategorized keyword and one sectioned keyword are both
    enabled. This proves it doesn't crash and reports "Uncategorized".
    """
    with migrated_schema() as conn:
        admin_service.create_keyword(conn, section=None, term_en="NEIR")
        admin_service.create_keyword(conn, section="telecom", term_en="BTRC")
        matching_id = _today_article(
            conn, url="https://x.test/nc1", url_norm="https://x.test/nc1", title="NEIR system launched nationwide"
        )
        conn.commit()

        result = admin_service.rematch_today(conn)
        assert result["matched"] == 1

        row = conn.execute("SELECT sections, keywords FROM articles WHERE id = %s", (matching_id,)).fetchone()
        assert row["sections"] == "Uncategorized"
        assert "NEIR" in row["keywords"]


def test_rematch_matches_bengali_only_keyword() -> None:
    """Regression for a matching.py bug this mission fixed:
    `keyword_rows_to_groups` used to require a truthy `term_en` before
    adding ANY Bengali synonym pattern, so a term_en="" / term_bn=<...>
    keyword silently compiled to zero patterns and could never match —
    even though creating one was already allowed (term_bn-only was
    always a valid `create_keyword` call). This proves it now matches.
    """
    with migrated_schema() as conn:
        admin_service.create_keyword(conn, section="telecom", term_en="", term_bn="স্পেকট্রাম")
        matching_id = _today_article(
            conn, url="https://x.test/bn1", url_norm="https://x.test/bn1", title="স্পেকট্রাম নিলাম শুরু হয়েছে"
        )
        conn.commit()

        result = admin_service.rematch_today(conn)
        assert result["matched"] == 1

        row = conn.execute("SELECT keywords FROM articles WHERE id = %s", (matching_id,)).fetchone()
        assert "স্পেকট্রাম" in row["keywords"]


# ── KEYWORDS-CONFIG: xlsx export/import ─────────────────────────────────


def test_export_then_reimport_xlsx_is_a_lossless_round_trip() -> None:
    with migrated_schema() as conn:
        admin_service.create_keyword(conn, section="telecom", term_en="spectrum", term_bn="স্পেকট্রাম")
        admin_service.create_keyword(conn, section="macro", term_en="reserves")
        admin_service.create_keyword(conn, section=None, term_en="", term_bn="ভোট")
        before = admin_service.list_keywords(conn)
        assert len(before) == 3

        content = admin_service.export_keywords_xlsx(conn)
        result = admin_service.import_keywords_xlsx(conn, content)

        # Every row matched by `id` (the export includes it) -> pure
        # updates, nothing created, nothing dropped, no duplicates.
        assert result == {"created": 0, "updated": 3, "skipped": 0, "errors": []}
        after = admin_service.list_keywords(conn)
        assert len(after) == 3
        assert {k["id"] for k in after} == {k["id"] for k in before}
        assert {(k["section"], k["term_en"], k["term_bn"]) for k in after} == {
            (k["section"], k["term_en"], k["term_bn"]) for k in before
        }


def test_import_xlsx_upserts_existing_row_by_id_and_inserts_new_row() -> None:
    with migrated_schema() as conn:
        existing = admin_service.create_keyword(conn, section="telecom", term_en="5G", term_bn="")

        wb = Workbook()
        ws = wb.active
        ws.append(["id", "section", "term_en", "term_bn", "enabled"])
        ws.append([existing["id"], "telecom", "5G", "৫জি", False])  # update via id
        ws.append(["", "macro", "reserves", "", True])  # brand new row, no id
        buf = io.BytesIO()
        wb.save(buf)

        result = admin_service.import_keywords_xlsx(conn, buf.getvalue())
        assert result["created"] == 1
        assert result["updated"] == 1
        assert result["skipped"] == 0
        assert result["errors"] == []

        updated = admin_service.get_keyword(conn, existing["id"])
        assert updated["term_bn"] == "৫জি"
        assert updated["enabled"] is False

        assert len(admin_service.list_keywords(conn)) == 2


def test_import_xlsx_matches_existing_row_by_triple_when_id_missing() -> None:
    with migrated_schema() as conn:
        existing = admin_service.create_keyword(conn, section="telecom", term_en="BTRC", term_bn="")

        wb = Workbook()
        ws = wb.active
        ws.append(["id", "section", "term_en", "term_bn", "enabled"])
        ws.append(["", "telecom", "BTRC", "", False])  # no id -- matched by (section, term_en, term_bn)
        buf = io.BytesIO()
        wb.save(buf)

        result = admin_service.import_keywords_xlsx(conn, buf.getvalue())
        assert result == {"created": 0, "updated": 1, "skipped": 0, "errors": []}
        assert admin_service.get_keyword(conn, existing["id"])["enabled"] is False
        assert len(admin_service.list_keywords(conn)) == 1


def test_import_xlsx_skips_rows_with_no_terms_and_reports_it() -> None:
    with migrated_schema() as conn:
        wb = Workbook()
        ws = wb.active
        ws.append(["id", "section", "term_en", "term_bn", "enabled"])
        ws.append(["", "telecom", "", "", True])  # both blank -- invalid
        ws.append(["", "telecom", "valid", "", True])
        buf = io.BytesIO()
        wb.save(buf)

        result = admin_service.import_keywords_xlsx(conn, buf.getvalue())
        assert result["created"] == 1
        assert result["skipped"] == 1
        assert len(result["errors"]) == 1
        assert len(admin_service.list_keywords(conn)) == 1


def test_import_xlsx_rejects_non_xlsx_content() -> None:
    with migrated_schema() as conn:
        with pytest.raises(ValidationError):
            admin_service.import_keywords_xlsx(conn, b"this is not a spreadsheet")


def test_import_xlsx_rejects_missing_required_columns() -> None:
    with migrated_schema() as conn:
        wb = Workbook()
        ws = wb.active
        ws.append(["section", "keyword"])  # wrong shape -- no term_en/term_bn
        ws.append(["telecom", "spectrum"])
        buf = io.BytesIO()
        wb.save(buf)
        with pytest.raises(ValidationError):
            admin_service.import_keywords_xlsx(conn, buf.getvalue())


# ── rematch_today ──────────────────────────────────────────────────────


def _today_article(conn, **overrides) -> int:
    fields = {
        "url": overrides.pop("url", "https://example.com/a"),
        "url_norm": overrides.pop("url_norm", "https://example.com/a"),
        "source": "Example",
        "title": overrides.pop("title", "BTRC fines Summit for bandwidth pricing"),
        "day": dt.date.today(),
        "review_status": overrides.pop("review_status", "pending"),
    }
    fields.update(overrides)
    article_id = insert_article(conn, **fields)
    assert article_id is not None
    return article_id


def test_rematch_today_matches_and_filters_pending_articles() -> None:
    with migrated_schema() as conn:
        admin_service.create_keyword(conn, section="telecom", term_en="BTRC")
        matching_id = _today_article(
            conn, url="https://x.test/1", url_norm="https://x.test/1", title="BTRC issues new rule"
        )
        non_matching_id = _today_article(
            conn, url="https://x.test/2", url_norm="https://x.test/2", title="Completely unrelated sports story"
        )
        conn.commit()

        result = admin_service.rematch_today(conn)
        assert result == {"matched": 1, "filtered_out": 1, "total": 2}

        matched_row = conn.execute(
            "SELECT review_status, keywords FROM articles WHERE id = %s", (matching_id,)
        ).fetchone()
        assert matched_row["review_status"] == "pending"
        assert "BTRC" in matched_row["keywords"]

        filtered_row = conn.execute("SELECT review_status FROM articles WHERE id = %s", (non_matching_id,)).fetchone()
        assert filtered_row["review_status"] == "filtered_out"


def test_rematch_today_never_touches_already_decided_articles() -> None:
    with migrated_schema() as conn:
        # No "BTRC" keyword configured at all -- if rematch touched this
        # approved row it would flip to filtered_out, which must not happen.
        approved_id = _today_article(
            conn, url="https://x.test/approved", url_norm="https://x.test/approved", review_status="approved"
        )
        conn.commit()

        result = admin_service.rematch_today(conn)
        assert result["total"] == 0

        row = conn.execute("SELECT review_status FROM articles WHERE id = %s", (approved_id,)).fetchone()
        assert row["review_status"] == "approved"


# ── Router-level ──────────────────────────────────────────────────────


def test_router_full_crud_and_rematch_flow_end_to_end(monkeypatch: pytest.MonkeyPatch) -> None:
    with migrated_schema() as conn:
        import app.api.routers.keywords as keywords_router_module

        monkeypatch.setattr(keywords_router_module, "get_connection", lambda: contextlib.nullcontext(conn))

        app = create_app()
        app.dependency_overrides[get_current_user] = lambda: ADMIN
        with TestClient(app) as client:
            created = client.post("/keywords", json={"section": "telecom", "term_en": "spectrum"})
            assert created.status_code == 201
            keyword_id = created.json()["id"]

            sections = client.get("/keywords/sections")
            assert sections.status_code == 200
            assert sections.json() == [{"section": "telecom", "count": 1}]

            toggled = client.post(f"/keywords/{keyword_id}/toggle")
            assert toggled.status_code == 200
            assert toggled.json()["enabled"] is False

            rematch = client.post("/keywords/rematch")
            assert rematch.status_code == 200
            assert rematch.json() == {"matched": 0, "filtered_out": 0, "total": 0}

            deleted = client.delete(f"/keywords/{keyword_id}")
            assert deleted.status_code == 204

            missing = client.patch(f"/keywords/{keyword_id}", json={"enabled": True})
            assert missing.status_code == 404


def test_router_create_without_section_and_rename_section(monkeypatch: pytest.MonkeyPatch) -> None:
    with migrated_schema() as conn:
        import app.api.routers.keywords as keywords_router_module

        monkeypatch.setattr(keywords_router_module, "get_connection", lambda: contextlib.nullcontext(conn))

        app = create_app()
        app.dependency_overrides[get_current_user] = lambda: ADMIN
        with TestClient(app) as client:
            created = client.post("/keywords", json={"term_en": "floating rate"})
            assert created.status_code == 201
            assert created.json()["section"] == "Uncategorized"

            created2 = client.post("/keywords", json={"section": "macro", "term_en": "reserves"})
            assert created2.status_code == 201

            renamed = client.patch("/keywords/sections/macro", json={"name": "Macroeconomy"})
            assert renamed.status_code == 200
            sections = {row["section"] for row in renamed.json()}
            assert "Macroeconomy" in sections
            assert "macro" not in sections


def test_router_export_xlsx_returns_binary_workbook(monkeypatch: pytest.MonkeyPatch) -> None:
    with migrated_schema() as conn:
        import app.api.routers.keywords as keywords_router_module

        monkeypatch.setattr(keywords_router_module, "get_connection", lambda: contextlib.nullcontext(conn))
        admin_service.create_keyword(conn, section="telecom", term_en="spectrum")

        app = create_app()
        app.dependency_overrides[get_current_user] = lambda: ADMIN
        with TestClient(app) as client:
            res = client.get("/keywords/export.xlsx")
            assert res.status_code == 200
            assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
            assert len(res.content) > 0


def test_router_import_xlsx_round_trip_via_raw_body(monkeypatch: pytest.MonkeyPatch) -> None:
    with migrated_schema() as conn:
        import app.api.routers.keywords as keywords_router_module

        monkeypatch.setattr(keywords_router_module, "get_connection", lambda: contextlib.nullcontext(conn))
        admin_service.create_keyword(conn, section="telecom", term_en="spectrum")
        admin_service.create_keyword(conn, section="macro", term_en="reserves")

        app = create_app()
        app.dependency_overrides[get_current_user] = lambda: ADMIN
        with TestClient(app) as client:
            exported = client.get("/keywords/export.xlsx")
            assert exported.status_code == 200

            imported = client.post(
                "/keywords/import",
                content=exported.content,
                headers={"Content-Type": "application/octet-stream"},
            )
            assert imported.status_code == 200
            body = imported.json()
            assert body == {"created": 0, "updated": 2, "skipped": 0, "errors": []}

            listed = client.get("/keywords", params={"section": "telecom"})
            assert len(listed.json()) == 1


def test_router_import_xlsx_bad_file_returns_422_not_500(monkeypatch: pytest.MonkeyPatch) -> None:
    with migrated_schema() as conn:
        import app.api.routers.keywords as keywords_router_module

        monkeypatch.setattr(keywords_router_module, "get_connection", lambda: contextlib.nullcontext(conn))

        app = create_app()
        app.dependency_overrides[get_current_user] = lambda: ADMIN
        with TestClient(app) as client:
            res = client.post(
                "/keywords/import", content=b"not an xlsx file", headers={"Content-Type": "application/octet-stream"}
            )
            assert res.status_code == 422
