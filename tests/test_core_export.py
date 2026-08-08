"""Tests for the Export builder (frontend-contract.md §3.4): GET
/export/preview, GET/POST /export/presets, POST /export.
"""

from __future__ import annotations

import contextlib
import itertools
from datetime import date, timedelta

import pytest
from fastapi.testclient import TestClient

from app.api.deps import AuthenticatedUser, get_current_user
from app.db.testing import migrated_schema, pg_available
from app.main import create_app

pytestmark = pytest.mark.skipif(not pg_available(), reason="no reachable Postgres (set TEST_DATABASE_URL)")

USER = AuthenticatedUser(id=1, email="analyst@grameenphone.com", role="analyst", display_name="Analyst", session_id=1)

_url_seq = itertools.count()


def _insert_article(conn, **fields):
    defaults = {
        "url": f"https://x.test/article-{next(_url_seq)}",
        "url_norm": "",
        "source": "Prothom Alo",
        "title": "Test headline",
        "day": date.today(),
        "sections": "",
        "keywords": "",
        "review_status": "pending",
    }
    defaults.update(fields)
    columns = list(defaults.keys())
    placeholders = ", ".join(f"%({c})s" for c in columns)
    row = conn.execute(
        f"INSERT INTO articles ({', '.join(columns)}) VALUES ({placeholders}) RETURNING id", defaults
    ).fetchone()
    return row["id"]


@pytest.fixture
def export_client(monkeypatch: pytest.MonkeyPatch):
    import app.api.routers.export as export_router_module

    with migrated_schema() as conn:
        monkeypatch.setattr(export_router_module, "get_connection", lambda: contextlib.nullcontext(conn))
        app = create_app()
        app.dependency_overrides[get_current_user] = lambda: USER
        with TestClient(app) as client:
            yield client, conn


def test_preview_unauthenticated_returns_401() -> None:
    app = create_app()
    with TestClient(app) as client:
        resp = client.get("/export/preview")
    assert resp.status_code == 401


def test_export_unauthenticated_returns_401() -> None:
    app = create_app()
    with TestClient(app) as client:
        resp = client.post("/export")
    assert resp.status_code == 401


def test_preview_counts_split_by_filter(export_client) -> None:
    client, conn = export_client
    today = date.today()
    _insert_article(conn, title="Approved", day=today, sections="Others", review_status="approved", llm_score=9)
    _insert_article(conn, title="Pending", day=today, sections="Others", review_status="pending", llm_score=7)
    _insert_article(conn, title="Filtered", day=today, sections="Others", review_status="filtered_out")
    conn.commit()

    resp = client.get("/export/preview", params={"filter": "approved", "scope": "today"})
    assert resp.status_code == 200
    data = resp.json()
    assert data["counts"]["approved"] == 1
    assert data["counts"]["matched"] == 2
    assert data["rows"] == 1
    assert "approved" in data["summary"]
    assert data["filename"].startswith("erm_approved_today_")


def test_preview_topn_caps_at_available_matched(export_client) -> None:
    client, conn = export_client
    today = date.today()
    for i in range(3):
        _insert_article(conn, title=f"M{i}", day=today, sections="Others", review_status="pending", llm_score=6 + i)
    conn.commit()

    resp = client.get("/export/preview", params={"filter": "topn", "topN": 25, "scope": "today"})
    assert resp.status_code == 200
    assert resp.json()["counts"]["topn"] == 3


def test_preview_invalid_filter_returns_422(export_client) -> None:
    client, _conn = export_client
    resp = client.get("/export/preview", params={"filter": "bogus"})
    assert resp.status_code == 422


def test_export_csv_returns_rows_for_scope(export_client) -> None:
    client, conn = export_client
    today = date.today()
    _insert_article(
        conn, title="Approved story", source="Daily Star", day=today,
        sections="Macroeconomy", review_status="approved", llm_score=8,
    )
    conn.commit()

    resp = client.post("/export", params={"filter": "approved", "scope": "today", "format": "csv"})
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("text/csv")
    assert "attachment" in resp.headers["content-disposition"]
    body = resp.text
    assert "Approved story" in body
    assert "Daily Star" in body


def test_export_xlsx_returns_a_real_workbook(export_client) -> None:
    client, conn = export_client
    _insert_article(
        conn, title="Xlsx story", source="Prothom Alo", day=date.today(),
        sections="Others", review_status="approved", llm_score=8,
    )
    conn.commit()

    resp = client.post("/export", params={"filter": "approved", "format": "xlsx"})
    assert resp.status_code == 200
    assert resp.headers["content-type"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert "attachment" in resp.headers["content-disposition"]

    import io

    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(resp.content))
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "Headline"
    assert any(row[0] == "Xlsx story" for row in rows[1:])


def test_export_columns_selection_and_order_are_honored(export_client) -> None:
    client, conn = export_client
    _insert_article(
        conn, title="Column test", source="Daily Star", day=date.today(),
        sections="Political", review_status="approved", llm_score=7,
    )
    conn.commit()

    resp = client.post("/export", params={"filter": "approved", "format": "csv", "columns": "source,headline"})
    assert resp.status_code == 200
    lines = resp.text.lstrip("﻿").strip().splitlines()
    assert lines[0].startswith("Source,Headline")
    assert lines[1].startswith("Daily Star,Column test")


def test_export_range_scope_filters_by_custom_dates(export_client) -> None:
    client, conn = export_client
    old_day = date.today() - timedelta(days=30)
    recent_day = date.today()
    _insert_article(conn, title="Old", day=old_day, sections="Others", review_status="approved", llm_score=8)
    _insert_article(conn, title="Recent", day=recent_day, sections="Others", review_status="approved", llm_score=8)
    conn.commit()

    resp = client.post(
        "/export",
        params={
            "filter": "approved",
            "scope": "range",
            "from": old_day.isoformat(),
            "to": old_day.isoformat(),
            "format": "csv",
        },
    )
    assert resp.status_code == 200
    assert "Old" in resp.text
    assert "Recent" not in resp.text


def test_presets_round_trip(export_client) -> None:
    client, _conn = export_client
    resp = client.post(
        "/export/presets",
        json={
            "name": "Monday digest",
            "note": "Top 50",
            "filter": "topn",
            "topN": 50,
            "scope": "week",
            "sections": "Macroeconomy,Political",
            "sort": "relevance",
            "format": "xlsx",
        },
    )
    assert resp.status_code == 201
    created = resp.json()
    assert created["name"] == "Monday digest"
    assert created["sections"] == ["Macroeconomy", "Political"]
    assert created["id"]

    listed = client.get("/export/presets").json()
    assert any(p["id"] == created["id"] and p["name"] == "Monday digest" for p in listed)


def test_create_preset_invalid_filter_returns_422(export_client) -> None:
    client, _conn = export_client
    resp = client.post("/export/presets", json={"name": "Bad", "filter": "nope"})
    assert resp.status_code == 422
