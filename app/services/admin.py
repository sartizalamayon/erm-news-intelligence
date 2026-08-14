"""Business logic for the admin console API (ADMIN-API mission).

Every function here takes an already-open `DictConnection` (the same
convention as `app.services.storage`/`app.services.auth`) and returns
plain, JSON-safe dicts (or dataclasses from a package this task consumes
rather than reimplements) — never a raw DB row or a datetime object — so
`app/api/routers/*.py` can hand results straight to a pydantic response
model. Route-layer admin enforcement (`require_admin`) is NOT this
module's job; every function here is reachable to any caller that holds
a connection, exactly like `app.services.storage`'s functions are. The
gate is `app.api.deps.require_admin`, applied per-route in the routers.

Sections below mirror the routers that call them:
  * Sources   -> app.api.routers.sources / .health (test, diagnose)
  * Keywords  -> app.api.routers.keywords
  * Prompts, provider routing/quota -> app.api.routers.prompts
  * Storage, export/import -> app.api.routers.dashboard
  * Source health, pipeline stage runs -> app.api.routers.health

Schema note: this task may not touch migrations/**. Two admin surfaces
have no dedicated table (prompt version history, per-stage-per-day
pipeline status) — both are stored as JSON in the generic `settings`
key/value table (0001_settings.sql), the exact pattern
`app.services.llm.routing`/`app.services.llm.quota` already established
for the identical constraint ("no dedicated quota table exists yet...
persisted through the existing generic settings key/value table").
"""

from __future__ import annotations

import copy
import io
import json
import re
import threading
from datetime import date, datetime, timedelta, timezone
from typing import Any

from openpyxl import Workbook, load_workbook
from psycopg.errors import UniqueViolation
from psycopg.types.json import Jsonb

from app.db.connection import DictConnection
from app.db.settings import get_setting, set_setting
from app.errors import SourceError, ValidationError
from app.logging import get_logger
from app.services.data_transfer import ExportManifest, ImportReport, export_to_jsonl, import_from_jsonl
from app.services.llm import quota as llm_quota
from app.services.llm import routing as llm_routing
from app.services.llm.prompts import (
    RELEVANCE_EXAMPLES_KEY,
    RELEVANCE_GUIDANCE_KEY,
    get_relevance_examples,
    get_relevance_guidance,
)
from app.services.matching import is_bengali, load_patterns, match_article
from app.services.storage import (
    PurgeResult,
    StorageUsage,
    get_storage_usage,
    purge_matched_articles,
    purge_stale_bodies,
    purge_stale_unmatched_articles,
)

_CHECK = "✓"
_CROSS = "✕"

# ═════════════════════════════════════════════════════════════════════
# Sources (frontend-contract.md §3.7 / §3.10)
# ═════════════════════════════════════════════════════════════════════

# `sources.method` is a single TEXT column (migrations/0002); the
# frontend contract's `s.methods[]` is an array (§3.7 notes this as an
# unresolved open question, "cannot be resolved from static reading").
# Storing multiple methods as a comma-joined string in that one column
# is the pragmatic choice available without a schema change — this task
# may not touch migrations/**.


logger = get_logger(__name__)


def _methods_to_str(methods: list[str]) -> str:
    return ",".join(m.strip() for m in methods if m.strip())


def _str_to_methods(method: str) -> list[str]:
    return [m.strip() for m in (method or "").split(",") if m.strip()]


def _source_row_to_dict(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"],
        "name": row["name"],
        "methods": _str_to_methods(row["method"]),
        "language": row["language"],
        "enabled": row["enabled"],
        "config": dict(row["config"]) if row.get("config") else {},
        "created_at": row["created_at"].isoformat() if row.get("created_at") else None,
        # Internal only — not part of SourceOut (app/api/schemas/admin.py),
        # pydantic silently drops unknown kwargs (default `extra="ignore"`).
        # Carried through so probe_source/diagnose_source below can key
        # off the stable registry correlation without a second query.
        # See migrations/0014_sources_registry_key.sql.
        "registry_key": row.get("registry_key"),
    }


def list_sources(conn: DictConnection, *, enabled: bool | None = None) -> list[dict[str, Any]]:
    if enabled is None:
        rows = conn.execute("SELECT * FROM sources ORDER BY name").fetchall()
    else:
        rows = conn.execute("SELECT * FROM sources WHERE enabled = %s ORDER BY name", (enabled,)).fetchall()
    return [_source_row_to_dict(dict(r)) for r in rows]


def get_source(conn: DictConnection, source_id: int) -> dict[str, Any] | None:
    row = conn.execute("SELECT * FROM sources WHERE id = %s", (source_id,)).fetchone()
    return _source_row_to_dict(dict(row)) if row else None


def create_source(
    conn: DictConnection,
    *,
    name: str,
    methods: list[str],
    language: str = "mixed",
    enabled: bool = True,
    config: dict[str, Any] | None = None,
) -> dict[str, Any]:
    if not name.strip():
        raise ValidationError("source name is required")
    if not methods:
        raise ValidationError("at least one fetch method is required")
    try:
        row = conn.execute(
            "INSERT INTO sources (name, method, config, language, enabled) VALUES (%s, %s, %s, %s, %s) RETURNING *",
            (name.strip(), _methods_to_str(methods), Jsonb(config or {}), language, enabled),
        ).fetchone()
    except UniqueViolation as exc:
        conn.rollback()
        raise ValidationError(f"a source named {name!r} already exists") from exc
    conn.commit()
    assert row is not None
    return _source_row_to_dict(dict(row))


def update_source(conn: DictConnection, source_id: int, **fields: Any) -> dict[str, Any] | None:
    existing = conn.execute("SELECT id FROM sources WHERE id = %s", (source_id,)).fetchone()
    if existing is None:
        return None

    sets: list[str] = []
    params: dict[str, Any] = {"id": source_id}
    if fields.get("name") is not None:
        sets.append("name = %(name)s")
        params["name"] = fields["name"]
    if fields.get("methods") is not None:
        if not fields["methods"]:
            raise ValidationError("at least one fetch method is required")
        sets.append("method = %(method)s")
        params["method"] = _methods_to_str(fields["methods"])
    if fields.get("language") is not None:
        sets.append("language = %(language)s")
        params["language"] = fields["language"]
    if fields.get("enabled") is not None:
        sets.append("enabled = %(enabled)s")
        params["enabled"] = fields["enabled"]
    if fields.get("config") is not None:
        sets.append("config = %(config)s")
        params["config"] = Jsonb(fields["config"])

    if not sets:
        return get_source(conn, source_id)

    try:
        row = conn.execute(f"UPDATE sources SET {', '.join(sets)} WHERE id = %(id)s RETURNING *", params).fetchone()
    except UniqueViolation as exc:
        conn.rollback()
        raise ValidationError(f"a source named {fields.get('name')!r} already exists") from exc
    conn.commit()
    assert row is not None
    return _source_row_to_dict(dict(row))


def delete_source(conn: DictConnection, source_id: int) -> bool:
    row = conn.execute("DELETE FROM sources WHERE id = %s RETURNING id", (source_id,)).fetchone()
    conn.commit()
    return row is not None


def toggle_source(conn: DictConnection, source_id: int) -> dict[str, Any] | None:
    row = conn.execute("UPDATE sources SET enabled = NOT enabled WHERE id = %s RETURNING *", (source_id,)).fetchone()
    if row is None:
        return None
    conn.commit()
    return _source_row_to_dict(dict(row))


# ── Live probes: test / diagnose, against the app.sources plugin layer ──
# Read-only use of app.sources.registry (P1-INGEST's territory, not
# edited here) — the registry maps a plugin `key` to a `NewsSource`
# class. This used to match the admin `sources` table's `name` column
# against that class's `display_name`, which broke the moment an admin
# renamed a source (PATCH /sources/{id} lets `name` drift freely, and
# the display-name match would then find nothing). Fixed: match on the
# stable `sources.registry_key` column instead (backfilled and kept in
# sync by migrations/0014_sources_registry_key.sql) — a rename no
# longer touches this lookup at all.

_registry_loaded = False


def _ensure_registry_loaded() -> None:
    global _registry_loaded
    if not _registry_loaded:
        from app.sources.registry import load_all

        load_all()
        _registry_loaded = True


def _plugin_for_registry_key(registry_key: str | None) -> Any | None:
    if not registry_key:
        return None
    from app.sources.registry import all_sources

    _ensure_registry_loaded()
    return all_sources().get(registry_key)


def _today_dhaka() -> date:
    # Dhaka is a fixed UTC+6 offset, no DST (matches
    # migrations/0003_articles.sql's `published` column comment).
    return (datetime.now(timezone.utc) + timedelta(hours=6)).date()


def probe_source(conn: DictConnection, source_id: int) -> dict[str, Any] | None:
    """`s.test` / `h.test` (frontend-contract.md §3.7/§3.10): a live
    "test this source now" call against the registered fetch plugin,
    returning a real result string, not a cached status read.
    """
    row = get_source(conn, source_id)
    if row is None:
        return None
    method_label = row["methods"][0] if row["methods"] else ""
    source_cls = _plugin_for_registry_key(row["registry_key"])
    if source_cls is None:
        return {
            "source_id": source_id,
            "ok": False,
            "fetched": 0,
            "today": 0,
            "method": method_label,
            "detail": f"no fetch plugin registered for {row['name']!r} yet (app.sources.registry)",
        }
    try:
        items = source_cls().iter_feed_items()
    except SourceError as exc:
        return {
            "source_id": source_id,
            "ok": False,
            "fetched": 0,
            "today": 0,
            "method": method_label,
            "detail": f"fetch failed: {exc}",
        }
    today = _today_dhaka()
    dhaka = timezone(timedelta(hours=6))
    today_count = sum(1 for it in items if it.published_at and it.published_at.astimezone(dhaka).date() == today)
    method_label = method_label or getattr(source_cls, "key", "")
    return {
        "source_id": source_id,
        "ok": True,
        "fetched": len(items),
        "today": today_count,
        "method": method_label,
        "detail": f"{len(items)} fetched, {today_count} from today (via {method_label})",
    }


def diagnose_source(conn: DictConnection, source_id: int) -> dict[str, Any] | None:
    """`h.diagnose` (frontend-contract.md §3.10) — genuine per-check
    diagnostics (reachable / feed discovery / one article body fetch /
    publish-date presence), never the mock's fabricated fixed string.
    """
    row = get_source(conn, source_id)
    if row is None:
        return None
    source_cls = _plugin_for_registry_key(row["registry_key"])
    if source_cls is None:
        return {
            "source_id": source_id,
            "reachable": False,
            "feed_discovery_ok": False,
            "article_url_found": False,
            "date_parse_ok": False,
            "detail": f"no fetch plugin registered for {row['name']!r} yet (app.sources.registry)",
        }

    instance = source_cls()
    try:
        items = instance.iter_feed_items()
    except SourceError as exc:
        return {
            "source_id": source_id,
            "reachable": False,
            "feed_discovery_ok": False,
            "article_url_found": False,
            "date_parse_ok": False,
            "detail": f"reachable {_CROSS} · {exc}",
        }

    feed_ok = len(items) > 0
    article_ok = False
    article_detail = ""
    if items:
        try:
            article = instance.fetch_article(items[0])
            article_ok = bool(article.body_text.strip())
        except SourceError as exc:
            article_detail = f" ({exc})"
    date_ok = any(it.published_at is not None for it in items)

    detail = " · ".join(
        [
            f"reachable {_CHECK}",
            f"feed discovery {_CHECK if feed_ok else _CROSS} ({len(items)} items)",
            f"article body fetch {_CHECK if article_ok else _CROSS}{article_detail}",
            f"date parse {_CHECK if date_ok else _CROSS + ' (no machine-readable publish date found)'}",
        ]
    )
    return {
        "source_id": source_id,
        "reachable": True,
        "feed_discovery_ok": feed_ok,
        "article_url_found": article_ok,
        "date_parse_ok": date_ok,
        "detail": detail,
    }


# ═════════════════════════════════════════════════════════════════════
# Keywords (frontend-contract.md §3.8; KEYWORDS-CONFIG: custom sections,
# nullable section, independently-optional bilingual terms, xlsx
# import/export)
# ═════════════════════════════════════════════════════════════════════

# `keywords.section` has never had a CHECK constraint or an enum at any
# layer (migrations/0002_reference_tables.sql: plain TEXT) -- an admin
# could already save an arbitrary custom section name before this
# mission touched anything. What this mission adds is (a) a NULL
# section (migrations/0012_keywords_section_nullable.sql) for "no
# section assigned", and (b) treating that NULL consistently as the
# display string below everywhere a section is read, written, or
# filtered on. "Uncategorized" is never itself stored as a real section
# value -- it is always a stand-in for SQL NULL, both coming out
# (_keyword_row_to_dict) and going in (_normalize_section), so a caller
# can round-trip it (e.g. via xlsx export -> re-import) without ever
# creating a second, indistinguishable section literally named that.
UNCATEGORIZED_SECTION = "Uncategorized"


def _normalize_section(value: str | None) -> str | None:
    """None / "" / whitespace-only / the literal "Uncategorized" all
    normalize to SQL NULL. Anything else is stored verbatim (stripped)
    -- this is the entire mechanism behind "custom sections": there is
    no allowlist to extend, any non-empty string an admin types becomes
    a valid section the moment a keyword is saved with it.
    """
    section = (value or "").strip()
    if not section or section == UNCATEGORIZED_SECTION:
        return None
    return section


def _keyword_row_to_dict(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "id": row["id"],
        "section": row["section"] or UNCATEGORIZED_SECTION,
        "term_en": row["term_en"],
        "term_bn": row["term_bn"],
        "enabled": row["enabled"],
        "created_at": row["created_at"].isoformat() if row.get("created_at") else None,
    }


def list_keywords(conn: DictConnection, *, section: str | None = None) -> list[dict[str, Any]]:
    if section:
        normalized = _normalize_section(section)
        if normalized is None:
            rows = conn.execute("SELECT * FROM keywords WHERE section IS NULL ORDER BY term_en").fetchall()
        else:
            rows = conn.execute(
                "SELECT * FROM keywords WHERE section = %s ORDER BY term_en", (normalized,)
            ).fetchall()
    else:
        rows = conn.execute("SELECT * FROM keywords ORDER BY section NULLS LAST, term_en").fetchall()
    return [_keyword_row_to_dict(dict(r)) for r in rows]


def list_keyword_sections(conn: DictConnection) -> list[dict[str, Any]]:
    """The list of sections "in use" (task 1's "browse/manage" surface)
    is deliberately DERIVED from distinct `keywords.section` values, not
    a separate authoritative table -- there is nothing a dedicated
    sections table would let an admin do that GROUP BY doesn't already
    give for free at this scale (a few dozen sections, at most), and a
    second table would just be one more place a section name could get
    out of sync with what keyword rows actually say.
    """
    rows = conn.execute(
        "SELECT COALESCE(section, %(u)s) AS section, count(*) AS count FROM keywords "
        "GROUP BY COALESCE(section, %(u)s) ORDER BY section",
        {"u": UNCATEGORIZED_SECTION},
    ).fetchall()
    return [{"section": r["section"], "count": r["count"]} for r in rows]


def rename_keyword_section(conn: DictConnection, section: str, new_name: str | None) -> int:
    """The "manage" half of task 1: since sections are derived values
    (no dedicated table -- see list_keyword_sections), "rename a
    section" IS "bulk-update every keyword row currently in it" -- there
    is nothing else to rename. Returns the number of rows touched.
    `new_name=None`/""/"Uncategorized" moves every matched row to no
    section at all, same normalization as create/update.
    """
    from_section = _normalize_section(section)
    to_section = _normalize_section(new_name)
    if from_section is None:
        rows = conn.execute(
            "UPDATE keywords SET section = %s WHERE section IS NULL RETURNING id", (to_section,)
        ).fetchall()
    else:
        rows = conn.execute(
            "UPDATE keywords SET section = %s WHERE section = %s RETURNING id", (to_section, from_section)
        ).fetchall()
    conn.commit()
    return len(rows)


def get_keyword(conn: DictConnection, keyword_id: int) -> dict[str, Any] | None:
    row = conn.execute("SELECT * FROM keywords WHERE id = %s", (keyword_id,)).fetchone()
    return _keyword_row_to_dict(dict(row)) if row else None


def create_keyword(
    conn: DictConnection, *, section: str | None = None, term_en: str = "", term_bn: str = "", enabled: bool = True
) -> dict[str, Any]:
    if not term_en.strip() and not term_bn.strip():
        raise ValidationError("at least one of term_en/term_bn is required")
    row = conn.execute(
        "INSERT INTO keywords (section, term_en, term_bn, enabled) VALUES (%s, %s, %s, %s) RETURNING *",
        (_normalize_section(section), term_en.strip(), term_bn.strip(), enabled),
    ).fetchone()
    conn.commit()
    assert row is not None
    return _keyword_row_to_dict(dict(row))


def update_keyword(conn: DictConnection, keyword_id: int, **fields: Any) -> dict[str, Any] | None:
    existing = conn.execute("SELECT id FROM keywords WHERE id = %s", (keyword_id,)).fetchone()
    if existing is None:
        return None
    sets: list[str] = []
    params: dict[str, Any] = {"id": keyword_id}
    # `section` is handled by presence-in-`fields`, not `is not None`:
    # unlike term_en/term_bn/enabled, `None` is a meaningful value here
    # (explicitly clear to "no section") that a caller must be able to
    # send, distinct from "field omitted, leave it alone".
    if "section" in fields:
        sets.append("section = %(section)s")
        params["section"] = _normalize_section(fields["section"])
    for col in ("term_en", "term_bn", "enabled"):
        if fields.get(col) is not None:
            sets.append(f"{col} = %({col})s")
            params[col] = fields[col]
    if not sets:
        return get_keyword(conn, keyword_id)
    row = conn.execute(f"UPDATE keywords SET {', '.join(sets)} WHERE id = %(id)s RETURNING *", params).fetchone()
    conn.commit()
    assert row is not None
    return _keyword_row_to_dict(dict(row))


def delete_keyword(conn: DictConnection, keyword_id: int) -> bool:
    row = conn.execute("DELETE FROM keywords WHERE id = %s RETURNING id", (keyword_id,)).fetchone()
    conn.commit()
    return row is not None


def toggle_keyword(conn: DictConnection, keyword_id: int) -> dict[str, Any] | None:
    row = conn.execute("UPDATE keywords SET enabled = NOT enabled WHERE id = %s RETURNING *", (keyword_id,)).fetchone()
    if row is None:
        return None
    conn.commit()
    return _keyword_row_to_dict(dict(row))


def rematch_today(conn: DictConnection) -> dict[str, int]:
    """`rematch` ("Rematch today", §3.8) — re-runs P2's keyword filter
    (app.services.matching) against today's already-fetched articles,
    using whatever the `keywords` table looks like RIGHT NOW (so an
    admin's just-added/edited/disabled keyword takes effect immediately
    on the next rematch, without waiting for tomorrow's collection run).

    Scoped to `review_status IN ('pending', 'filtered_out')` only —
    deliberately never touches 'approved'/'rejected' rows, so a rematch
    can never silently undo an analyst's already-made review decision.
    """
    patterns = load_patterns(conn)
    rows = conn.execute(
        "SELECT id, title FROM articles WHERE day = CURRENT_DATE AND review_status IN ('pending', 'filtered_out')"
    ).fetchall()
    matched = filtered_out = 0
    for row in rows:
        result = match_article({"title": row["title"]}, patterns)
        conn.execute(
            "UPDATE articles SET sections = %s, keywords = %s, review_status = %s, match_type = %s WHERE id = %s",
            (result["sections"], result["keywords"], result["review_status"], result["match_type"], row["id"]),
        )
        if result["review_status"] == "pending":
            matched += 1
        else:
            filtered_out += 1
    conn.commit()
    return {"matched": matched, "filtered_out": filtered_out, "total": len(rows)}


def reset_day(conn: DictConnection, *, day: date | None = None, keep_reviewed: bool = True) -> dict[str, Any]:
    """Make a day collectable again from scratch — the supported way to
    re-run a collection you've already run.

    This exists because "just DELETE the rows" is a trap. `seen_urls` is
    a permanent tombstone table (migrations/0016): the
    `articles_remember_deleted_urls` trigger records every deleted
    url_norm precisely so a purged article can never silently reappear.
    That is correct for the retention purge and catastrophic for a
    deliberate re-run — the day's articles vanish, their URLs stay
    remembered, and every later fetch reports "0 new" for sources that
    are working perfectly. That exact failure was observed live
    (2026-08-14: 8,931 orphaned tombstones, sources reporting "178
    fetched · 4 new"), which is why the cleanup is done HERE, in the
    same transaction as the delete, rather than left to each caller.

    `keep_reviewed=True` (the default) refuses to discard rows an
    analyst has already approved or rejected — re-running a collection
    is a data-collection operation and must never destroy review work.
    Pass False only for a genuine full wipe of the day.
    """
    target = day or _today_dhaka()
    status_filter = "" if not keep_reviewed else " AND review_status IN ('pending', 'filtered_out')"

    with conn.transaction():
        deleted = conn.execute(
            f"DELETE FROM articles WHERE day = %s{status_filter} RETURNING url_norm",
            (target,),
        ).fetchall()
        norms = [row["url_norm"] for row in deleted if row["url_norm"]]
        # The trigger above has already written a tombstone for each of
        # these. Drop exactly those, and only those, so the URLs become
        # re-collectable — without touching tombstones from real purges.
        forgotten = 0
        if norms:
            forgotten = len(
                conn.execute(
                    "DELETE FROM seen_urls WHERE url_norm = ANY(%s) RETURNING url_norm",
                    (norms,),
                ).fetchall()
            )
    return {"deleted": len(deleted), "forgotten_urls": forgotten, "day": target.isoformat()}


# ── Excel export/import (KEYWORDS-CONFIG tasks 4/5) ─────────────────────
#
# Both writer and reader are `openpyxl` (already a dependency —
# app.services.export uses it identically for GET /export). Column shape
# is fixed and shared between the two directions: `id, section, term_en,
# term_bn, enabled`. `id` is included in the export specifically so a
# re-import can match by identity rather than only by content — see
# import_keywords_xlsx's docstring for why that matters.

KEYWORD_EXPORT_COLUMNS = ("id", "section", "term_en", "term_bn", "enabled")


def export_keywords_xlsx(conn: DictConnection) -> bytes:
    """Every keyword, one row each, in `KEYWORD_EXPORT_COLUMNS` order.
    `section` is written as the display string "Uncategorized" for a
    NULL-section row (`_keyword_row_to_dict`'s coercion) — so re-running
    this same file through `import_keywords_xlsx` round-trips it back to
    NULL via that function's own `_normalize_section` call, with no
    special-casing needed on either side.
    """
    rows = conn.execute("SELECT * FROM keywords ORDER BY section NULLS LAST, term_en").fetchall()
    wb = Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = "Keywords"
    ws.append(list(KEYWORD_EXPORT_COLUMNS))
    for r in rows:
        d = _keyword_row_to_dict(dict(r))
        ws.append([d["id"], d["section"], d["term_en"], d["term_bn"], d["enabled"]])
    for i in range(1, len(KEYWORD_EXPORT_COLUMNS) + 1):
        ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = 22
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _parse_bool_cell(value: Any) -> bool:
    """openpyxl hands back a real `bool` for an actual Excel boolean
    cell; anything typed in as text (a re-saved/hand-edited sheet) is
    tolerant-parsed — only an explicit false-ish token turns a row off,
    everything else (including a blank cell) defaults enabled=True,
    matching `KeywordCreate`'s own default.
    """
    if isinstance(value, bool):
        return value
    if value is None:
        return True
    return str(value).strip().lower() not in ("false", "0", "no", "n", "off")


def import_keywords_xlsx(conn: DictConnection, content: bytes) -> dict[str, Any]:
    """UPSERT semantics (task 5's "auto-update"): a row is matched
    first by `id` (if the sheet has one and it names a keyword that
    still exists), then by the `(section, term_en, term_bn)` triple —
    either match UPDATEs that row in place; no match INSERTs a new one.
    Nothing already in the table is ever deleted just for being absent
    from the sheet — a partial or filtered export re-imported cleanly
    can't wipe the rest of the keyword list.

    Row-level errors (bad id, unexpected exception) are caught per-row
    and reported in `errors` rather than aborting the whole import or
    raising a 500 — only a structurally bad file (wrong columns, not a
    real .xlsx) raises `ValidationError` (-> 422) for the request as a
    whole.
    """
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - any parse failure is a bad-file 422, not a 500
        raise ValidationError(f"not a valid .xlsx file: {exc}") from exc

    ws = wb.active
    if ws is None:
        raise ValidationError("workbook has no sheets")

    rows_iter = ws.iter_rows(values_only=True)
    try:
        header = next(rows_iter)
    except StopIteration:
        raise ValidationError("workbook is empty (no header row)") from None

    header_norm = [str(h).strip().lower() if h is not None else "" for h in header]
    required = {"term_en", "term_bn"}
    if not required.issubset(header_norm):
        raise ValidationError(
            f"expected columns {list(KEYWORD_EXPORT_COLUMNS)} (found {header_norm!r})"
        )
    col_idx = {name: idx for idx, name in enumerate(header_norm) if name}

    def cell(row: tuple[Any, ...], name: str) -> Any:
        idx = col_idx.get(name)
        if idx is None or idx >= len(row):
            return None
        return row[idx]

    created = updated = skipped = 0
    errors: list[str] = []

    for i, row in enumerate(rows_iter, start=2):  # row 1 is the header
        if row is None or all(v is None for v in row):
            continue
        try:
            raw_id = cell(row, "id")
            section = str(cell(row, "section") or "").strip()
            term_en = str(cell(row, "term_en") or "").strip()
            term_bn = str(cell(row, "term_bn") or "").strip()
            enabled = _parse_bool_cell(cell(row, "enabled"))

            if not term_en and not term_bn:
                skipped += 1
                errors.append(f"row {i}: skipped — both term_en and term_bn are blank")
                continue

            normalized_section = _normalize_section(section)

            existing_id: int | None = None
            if raw_id not in (None, ""):
                try:
                    candidate = get_keyword(conn, int(raw_id))
                except (TypeError, ValueError):
                    candidate = None
                if candidate is not None:
                    existing_id = candidate["id"]

            if existing_id is None:
                if normalized_section is None:
                    match = conn.execute(
                        "SELECT id FROM keywords WHERE section IS NULL AND term_en = %s AND term_bn = %s",
                        (term_en, term_bn),
                    ).fetchone()
                else:
                    match = conn.execute(
                        "SELECT id FROM keywords WHERE section = %s AND term_en = %s AND term_bn = %s",
                        (normalized_section, term_en, term_bn),
                    ).fetchone()
                existing_id = match["id"] if match else None

            if existing_id is not None:
                conn.execute(
                    "UPDATE keywords SET section = %s, term_en = %s, term_bn = %s, enabled = %s WHERE id = %s",
                    (normalized_section, term_en, term_bn, enabled, existing_id),
                )
                updated += 1
            else:
                conn.execute(
                    "INSERT INTO keywords (section, term_en, term_bn, enabled) VALUES (%s, %s, %s, %s)",
                    (normalized_section, term_en, term_bn, enabled),
                )
                created += 1
        except Exception as exc:  # noqa: BLE001 - one bad row must not fail the whole import
            skipped += 1
            errors.append(f"row {i}: {exc}")

    conn.commit()
    return {"created": created, "updated": updated, "skipped": skipped, "errors": errors}


# ═════════════════════════════════════════════════════════════════════
# AI prompts + versioning (frontend-contract.md §3.9)
# ═════════════════════════════════════════════════════════════════════
#
# No dedicated table (this task may not touch migrations/**) — version
# history is JSON in the `settings` KV table, keyed
# "prompt_versions:{field}", the identical pattern
# app.services.llm.routing/quota already use for the same constraint.
# The ACTIVE version's content is ALSO mirrored into the plain
# `relevance_guidance`/`relevance_examples` settings keys
# (app.services.llm.prompts' RELEVANCE_GUIDANCE_KEY/RELEVANCE_EXAMPLES_KEY)
# on every save/activate, since those are the keys P3's actual runtime
# read path (get_relevance_guidance/get_relevance_examples) consults —
# versioning must never require a change to that frozen read path.

PROMPT_FIELDS: dict[str, str] = {
    "guidance": RELEVANCE_GUIDANCE_KEY,
    "examples": RELEVANCE_EXAMPLES_KEY,
}


def _validate_prompt_field(field: str) -> None:
    if field not in PROMPT_FIELDS:
        raise ValidationError(f"unknown prompt field {field!r} — must be one of {sorted(PROMPT_FIELDS)}")


def _versions_settings_key(field: str) -> str:
    return f"prompt_versions:{field}"


def _load_prompt_versions(conn: DictConnection, field: str) -> dict[str, Any]:
    raw = get_setting(conn, _versions_settings_key(field), default="")
    if raw.strip():
        try:
            data = json.loads(raw)
            if isinstance(data, dict) and data.get("versions"):
                return data
        except json.JSONDecodeError:
            pass
    # No history yet — seed version 1 from whatever P3's runtime read
    # path would currently return (get_relevance_guidance/_examples's own
    # fallback semantics), so there is always at least one real,
    # retrievable version rather than an empty list.
    current = get_relevance_guidance(conn) if field == "guidance" else get_relevance_examples(conn)
    seed = {
        "active": 1,
        "versions": [
            {
                "version": 1,
                "content": current,
                "created_at": datetime.now(timezone.utc).isoformat(),
                "created_by": "system",
                "note": "initial version, seeded from the existing configured value",
            }
        ],
    }
    set_setting(conn, _versions_settings_key(field), json.dumps(seed))
    return seed


def _save_prompt_versions(conn: DictConnection, field: str, data: dict[str, Any]) -> None:
    set_setting(conn, _versions_settings_key(field), json.dumps(data))


def list_prompt_versions(conn: DictConnection, field: str) -> list[dict[str, Any]]:
    _validate_prompt_field(field)
    data = _load_prompt_versions(conn, field)
    active = data["active"]
    return [{**v, "active": v["version"] == active} for v in data["versions"]]


def save_prompt_content(
    conn: DictConnection, field: str, content: str, *, actor_email: str, note: str = ""
) -> dict[str, Any]:
    """Every save creates a NEW version and makes it active immediately
    (DoD #3: "every change creates a new version"). Never mutates an
    existing version's content — history is append-only.
    """
    _validate_prompt_field(field)
    data = _load_prompt_versions(conn, field)
    next_version = max((v["version"] for v in data["versions"]), default=0) + 1
    entry = {
        "version": next_version,
        "content": content,
        "created_at": datetime.now(timezone.utc).isoformat(),
        "created_by": actor_email,
        "note": note,
    }
    data["versions"].append(entry)
    data["active"] = next_version
    _save_prompt_versions(conn, field, data)
    set_setting(conn, PROMPT_FIELDS[field], content)
    conn.commit()
    return {**entry, "active": True}


def activate_prompt_version(conn: DictConnection, field: str, version: int) -> dict[str, Any] | None:
    """Repoints "active" at an EXISTING version (DoD #3: "the active
    version is selectable") — does not create a new version entry, since
    the content itself hasn't changed, only which one is live.
    """
    _validate_prompt_field(field)
    data = _load_prompt_versions(conn, field)
    match = next((v for v in data["versions"] if v["version"] == version), None)
    if match is None:
        return None
    data["active"] = version
    _save_prompt_versions(conn, field, data)
    set_setting(conn, PROMPT_FIELDS[field], match["content"])
    conn.commit()
    return {**match, "active": True}


def get_prompt_config(conn: DictConnection) -> dict[str, Any]:
    guidance_data = _load_prompt_versions(conn, "guidance")
    examples_data = _load_prompt_versions(conn, "examples")
    guidance_active = next(v for v in guidance_data["versions"] if v["version"] == guidance_data["active"])
    examples_active = next(v for v in examples_data["versions"] if v["version"] == examples_data["active"])
    approved_row = conn.execute("SELECT count(*) AS n FROM articles WHERE review_status = 'approved'").fetchone()
    rejected_row = conn.execute("SELECT count(*) AS n FROM articles WHERE review_status = 'rejected'").fetchone()
    return {
        "guidance": guidance_active["content"],
        "examples": examples_active["content"],
        "guidance_version": guidance_data["active"],
        "examples_version": examples_data["active"],
        "approved_count": approved_row["n"] if approved_row else 0,
        "rejected_count": rejected_row["n"] if rejected_row else 0,
    }


# ── Keyword-candidate mining (§3.9 candidates[]) ────────────────────────
# A lean, dependency-free word-frequency co-occurrence miner over real
# approve/reject decision history — no new dependency (GATE 0 / mission
# "LEAN"), replacing the mock's static illustrative sample.

_STOPWORDS = {
    "the",
    "a",
    "an",
    "to",
    "of",
    "in",
    "on",
    "for",
    "and",
    "or",
    "is",
    "are",
    "was",
    "were",
    "by",
    "with",
    "as",
    "at",
    "from",
    "after",
    "over",
    "new",
    "says",
    "amid",
    "into",
    "than",
    "will",
    "has",
    "have",
    "had",
    "its",
    "his",
    "her",
    "their",
    "this",
    "that",
    "not",
    "no",
}


def _is_low_information(word: str) -> bool:
    """Reject tokens that carry no real editorial signal as a keyword:
    pure numbers ("100", "2026") and tokens where digits are the
    majority (a stray "$100" surviving tokenization as "100"). Strictly
    a MAJORITY, not merely half, so a meaningful mixed alphanumeric
    term where letters still dominate (e.g. "4glte") is not rejected
    just for containing a digit.
    """
    if not word:
        return True
    digits = sum(1 for ch in word if ch.isdigit())
    return digits / len(word) > 0.5


def _tokenize(text: str | None) -> set[str]:
    words = re.findall(r"[\w']+", (text or "").lower())
    return {w for w in words if len(w) > 2 and w not in _STOPWORDS and not _is_low_information(w)}


def mine_keyword_candidates(
    conn: DictConnection, *, limit: int = 6, sample: int = 300, min_count: int = 2
) -> list[dict[str, Any]]:
    approved_rows = conn.execute(
        "SELECT title, title_en FROM articles WHERE review_status = 'approved' ORDER BY id DESC LIMIT %s",
        (sample,),
    ).fetchall()
    rejected_rows = conn.execute(
        "SELECT title, title_en FROM articles WHERE review_status = 'rejected' ORDER BY id DESC LIMIT %s",
        (sample,),
    ).fetchall()
    existing_terms = {
        (row["term_en"] or "").strip().lower() for row in conn.execute("SELECT term_en FROM keywords").fetchall()
    }

    approved_counts: dict[str, int] = {}
    rejected_counts: dict[str, int] = {}
    for row in approved_rows:
        for word in _tokenize(row["title_en"] or row["title"]):
            approved_counts[word] = approved_counts.get(word, 0) + 1
    for row in rejected_rows:
        for word in _tokenize(row["title_en"] or row["title"]):
            rejected_counts[word] = rejected_counts.get(word, 0) + 1

    candidates = [
        {
            "word": word,
            "lang": "BN" if is_bengali(word) else "EN",
            "approved": count,
            "rejected": rejected_counts.get(word, 0),
        }
        for word, count in approved_counts.items()
        # A word seen once is weak evidence for a keyword suggestion —
        # require it to recur across at least `min_count` approved
        # articles before surfacing it as a candidate.
        if word not in existing_terms and count >= min_count
    ]
    candidates.sort(key=lambda c: (-int(c["approved"]), int(c["rejected"])))
    return candidates[:limit]


def _recent_titles(conn: DictConnection, review_status: str, limit: int = 40) -> list[str]:
    rows = conn.execute(
        "SELECT title_en, title FROM articles WHERE review_status = %s ORDER BY id DESC LIMIT %s",
        (review_status, limit),
    ).fetchall()
    return [(row["title_en"] or row["title"]) for row in rows]


def _build_client_for_feature(conn: DictConnection, feature: str) -> Any:
    """Lazy import of app.services.llm.registry — deliberately NOT a
    module-level import. That module imports app.services.llm.transport's
    concrete transport classes, and both packages are P3-LLM's actively
    concurrent territory; a module-level import here would make this
    entire file (and therefore every admin route, including plain
    keyword/source CRUD that has nothing to do with the LLM layer)
    unimportable during any moment P3-LLM's package is mid-refactor.
    Deferred to call time, the blast radius of that package being
    transiently broken is contained to the one action that genuinely
    needs it (suggest_refinement) rather than the whole admin API.
    """
    from app.services.llm.registry import build_client_for_feature

    return build_client_for_feature(conn, feature)


def suggest_refinement(conn: DictConnection) -> str:
    """ "Suggest a refinement" (§3.9's `suggest`) — a genuine LLM call
    server-side, built on top of P3-LLM's already-provided pieces:
    `LLMClient.suggest_guidance()` (app.services.llm.provider) and its
    routing slot, folded into the `summarization` feature
    (app.services.llm.routing's own docstring names this exact reason).
    Never auto-applied — the caller (PUT /prompts/guidance) is a
    separate, explicit save.
    """
    approved = _recent_titles(conn, "approved")
    rejected = _recent_titles(conn, "rejected")
    client = _build_client_for_feature(conn, llm_routing.SUMMARIZATION)
    return client.suggest_guidance(approved, rejected)


# ═════════════════════════════════════════════════════════════════════
# Provider routing / quota config (ADR §9.1, DoD #4)
# ═════════════════════════════════════════════════════════════════════


def _provider_choice_to_dict(choice: llm_routing.ProviderChoice | None) -> dict[str, Any] | None:
    if choice is None:
        return None
    return {"provider": choice.provider, "model": choice.model}


def _feature_route_to_dict(route: llm_routing.FeatureRoute) -> dict[str, Any]:
    return {"primary": _provider_choice_to_dict(route.primary), "secondary": _provider_choice_to_dict(route.secondary)}


def get_routing_config(conn: DictConnection) -> dict[str, Any]:
    config = llm_routing.load_routing_config(conn)
    return {feature: _feature_route_to_dict(route) for feature, route in config.routes.items()}


def save_routing_config(conn: DictConnection, routes: dict[str, Any]) -> dict[str, Any]:
    """Read-modify-write over the FULL 5-feature map: an admin changing
    just one feature's route (the common case — "route scoring to
    Gemini") should not have to resend the other four untouched. Starts
    from `load_routing_config()`'s current state (already-defaulted for
    any feature the admin has never touched) and overlays only the
    feature(s) present in `routes`, then persists the complete map via
    `app.services.llm.routing.save_routing_config`, which itself always
    writes the full 5-feature map (see that function's own docstring).
    """
    current = llm_routing.load_routing_config(conn)
    built: dict[str, llm_routing.FeatureRoute] = dict(current.routes)
    for feature, route in routes.items():
        primary_raw = route["primary"]
        primary = llm_routing.ProviderChoice(provider=primary_raw["provider"], model=primary_raw.get("model"))
        secondary_raw = route.get("secondary")
        secondary = (
            llm_routing.ProviderChoice(provider=secondary_raw["provider"], model=secondary_raw.get("model"))
            if secondary_raw
            else None
        )
        built[feature] = llm_routing.FeatureRoute(primary=primary, secondary=secondary)
    config = llm_routing.RoutingConfig(routes=built)
    llm_routing.save_routing_config(conn, config)
    conn.commit()
    return {feature: _feature_route_to_dict(route) for feature, route in config.routes.items()}


def _provider_config_row_to_dict(row: dict[str, Any]) -> dict[str, Any]:
    return {
        "provider": row["provider"],
        "model": row["model"],
        "role": row["role"],
        "api_key_last4": row["api_key_last4"],
        "has_key": bool(row["has_key"]),
        "updated_at": row["updated_at"].isoformat() if row.get("updated_at") else None,
    }


# Providers `provider_configs` (ARCH-DB's table) can represent — mirrors
# that table's own CHECK constraint and
# app.services.llm.registry.PROVIDER_CONFIGS_PROVIDERS (not imported
# from there directly: that module transitively imports
# app.services.llm.transport's concrete transport classes and is
# P3-LLM's actively concurrent territory — see
# _build_client_for_feature's docstring for why this file never imports
# it at module level). `llm_routing.KNOWN_PROVIDERS` already holds the
# identical tuple for the identical reason and IS safe to import here
# (app.services.llm.routing has no such transitive dependency).
_PROVIDER_CONFIGS_PROVIDERS = llm_routing.KNOWN_PROVIDERS


def list_provider_configs(conn: DictConnection) -> list[dict[str, Any]]:
    rows = conn.execute(
        "SELECT provider, model, role, api_key_last4, updated_at, "
        "(api_key_ciphertext IS NOT NULL) AS has_key FROM provider_configs ORDER BY provider"
    ).fetchall()
    return [_provider_config_row_to_dict(dict(r)) for r in rows]


_PROVIDER_ROLES = ("primary", "secondary", "disabled")


def upsert_provider_config(
    conn: DictConnection,
    provider: str,
    *,
    model: str | None = None,
    role: str | None = None,
    api_key: str | None = None,
    updated_by: int | None = None,
) -> dict[str, Any]:
    """Writes to `provider_configs` (ARCH-DB's table) — restricted to
    `groq`/`gemini` because that table's own `CHECK (provider IN ('groq',
    'gemini'))` constraint makes it structurally impossible to hold
    anything else (a migration this task may not add). This is
    deliberately NOT the generic "admin adds a new provider" surface —
    for that (an arbitrary provider name, its own base_url/adapter, and
    one or more pooled API keys), see `save_provider_definition` /
    `save_provider_credential` below, backed by
    `app.services.llm.provider_registry` (P3-LLM's Task 3 addition).

    Never accepts or returns a plaintext key beyond this one write
    (ADR §9.1) — `api_key`, if given, is encrypted immediately via
    `app.db.crypto` and only the ciphertext + masked suffix are stored;
    the return value never includes it (see ProviderConfigOut).
    """
    if provider not in _PROVIDER_CONFIGS_PROVIDERS:
        raise ValidationError(f"unknown provider {provider!r} — must be one of {_PROVIDER_CONFIGS_PROVIDERS}")
    if role is not None and role not in _PROVIDER_ROLES:
        raise ValidationError(f"role must be one of {_PROVIDER_ROLES}, got {role!r}")

    existing = conn.execute("SELECT model, role FROM provider_configs WHERE provider = %s", (provider,)).fetchone()
    resolved_model = model if model is not None else (existing["model"] if existing else "")
    resolved_role = role if role is not None else (existing["role"] if existing else "secondary")

    ciphertext = None
    last4 = None
    if api_key:
        from app.db.crypto import encrypt_api_key

        ciphertext, last4 = encrypt_api_key(api_key)

    if existing is None:
        conn.execute(
            "INSERT INTO provider_configs (provider, model, role, api_key_ciphertext, api_key_last4, updated_by) "
            "VALUES (%s, %s, %s, %s, %s, %s)",
            (provider, resolved_model, resolved_role, ciphertext, last4 or "", updated_by),
        )
    elif ciphertext is not None:
        conn.execute(
            "UPDATE provider_configs SET model = %s, role = %s, api_key_ciphertext = %s, api_key_last4 = %s, "
            "updated_by = %s, updated_at = now() WHERE provider = %s",
            (resolved_model, resolved_role, ciphertext, last4, updated_by, provider),
        )
    else:
        conn.execute(
            "UPDATE provider_configs SET model = %s, role = %s, updated_by = %s, updated_at = now() "
            "WHERE provider = %s",
            (resolved_model, resolved_role, updated_by, provider),
        )
    conn.commit()
    row = conn.execute(
        "SELECT provider, model, role, api_key_last4, updated_at, "
        "(api_key_ciphertext IS NOT NULL) AS has_key FROM provider_configs WHERE provider = %s",
        (provider,),
    ).fetchone()
    assert row is not None
    return _provider_config_row_to_dict(dict(row))


def _quota_status_to_dict(status: llm_quota.QuotaStatus) -> dict[str, Any]:
    return {
        "provider": status.provider,
        "day": status.day.isoformat(),
        "tokens_used": status.used.tokens,
        "requests_used": status.used.requests,
        "tpd": status.limits.tpd,
        "rpd": status.limits.rpd,
        "tokens_remaining": status.tokens_remaining,
        "requests_remaining": status.requests_remaining,
        "tokens_used_pct": status.tokens_used_pct,
        "requests_used_pct": status.requests_used_pct,
        "exhausted": status.exhausted,
    }


def get_quota_status(conn: DictConnection) -> list[dict[str, Any]]:
    """Every provider worth showing remaining-budget for: the two ADR
    §9.1-ruled ones, every admin-defined provider
    (app.services.llm.provider_registry — mistral, or anything an admin
    has added), and any provider that already has a saved quota limit
    even if its definition was later removed. Never a hardcoded 2-vendor
    list — provider identities are runtime config now (Task 3).
    """
    limits = llm_quota.load_provider_limits(conn)
    definitions = load_provider_definitions(conn)
    providers = sorted(set(llm_routing.KNOWN_PROVIDERS) | set(definitions.keys()) | set(limits.keys()))
    tracker = llm_quota.QuotaTracker(store=llm_quota.SettingsQuotaStore(conn=conn), limits=limits)
    return [_quota_status_to_dict(tracker.status(provider)) for provider in providers]


def save_quota_limit(conn: DictConnection, provider: str, *, tpd: int | None, rpd: int | None) -> dict[str, Any]:
    if not provider.strip():
        raise ValidationError("provider must be non-empty")
    limits = llm_quota.load_provider_limits(conn)
    limits[provider] = llm_quota.ProviderLimits(tpd=tpd, rpd=rpd)
    llm_quota.save_provider_limits(conn, limits)
    conn.commit()
    tracker = llm_quota.QuotaTracker(store=llm_quota.SettingsQuotaStore(conn=conn), limits=limits)
    return _quota_status_to_dict(tracker.status(provider))


# ── Generic provider registry: arbitrary providers + pooled credentials
# (ADR §9.1 / owner's Task 3: "admins add API keys for arbitrary
# providers from the dashboard") ─────────────────────────────────────
#
# app.services.llm.provider_registry is imported lazily (function-local),
# same reasoning as _build_client_for_feature above: it's P3-LLM's
# concurrent territory, and this file must stay importable (so plain
# keyword/source CRUD keeps working) even if that package is transiently
# broken mid-edit.


def _definition_to_dict(definition: Any) -> dict[str, Any]:
    return {
        "provider": definition.provider,
        "adapter": definition.adapter,
        "base_url": definition.base_url,
        "default_model": definition.default_model,
        "enabled": definition.enabled,
        "note": definition.note,
    }


def load_provider_definitions(conn: DictConnection) -> dict[str, Any]:
    from app.services.llm.provider_registry import load_provider_definitions as _load

    return _load(conn)


def list_provider_definitions(conn: DictConnection) -> list[dict[str, Any]]:
    definitions = load_provider_definitions(conn)
    return [_definition_to_dict(d) for _, d in sorted(definitions.items())]


def save_provider_definition(
    conn: DictConnection,
    provider: str,
    *,
    adapter: str,
    base_url: str | None,
    default_model: str,
    enabled: bool = True,
    note: str = "",
) -> dict[str, Any]:
    """Admin write path for "add a provider" — a brand-new provider name
    with its own adapter kind (`openai_compatible` needs `base_url`;
    `gemini` does not), default model, and enabled flag. No code change,
    no migration — this is the entire mechanism Task 3 asks for.
    """
    if not provider.strip():
        raise ValidationError("provider must be non-empty")
    from app.services.llm.provider_registry import ProviderDefinition
    from app.services.llm.provider_registry import save_provider_definitions as _save

    definitions = load_provider_definitions(conn)
    definitions[provider] = ProviderDefinition(
        provider=provider, adapter=adapter, base_url=base_url, default_model=default_model, enabled=enabled, note=note
    )  # raises ValidationError for an unknown adapter / missing base_url
    _save(conn, definitions)
    conn.commit()
    return _definition_to_dict(definitions[provider])


def list_provider_credentials(conn: DictConnection, provider: str) -> list[dict[str, Any]]:
    """Masked suffix only (ADR §9.1) — `load_provider_credentials`
    returns DECRYPTED keys for the LLM client's own use; this admin-facing
    listing must never pass that plaintext further than this function.
    """
    from app.logging import mask
    from app.services.llm.provider_registry import load_provider_credentials as _load

    return [
        {"provider": c.provider, "key_id": c.key_id, "api_key_last4": mask(c.api_key)} for c in _load(conn, provider)
    ]


def save_provider_credential(conn: DictConnection, provider: str, key_id: str, api_key: str) -> dict[str, Any]:
    from app.logging import mask
    from app.services.llm.provider_registry import save_provider_credential as _save

    _save(conn, provider, key_id, api_key)  # raises ValidationError if key_id is blank
    conn.commit()
    return {"provider": provider, "key_id": key_id, "api_key_last4": mask(api_key)}


def remove_provider_credential(conn: DictConnection, provider: str, key_id: str) -> None:
    from app.services.llm.provider_registry import remove_provider_credential as _remove

    _remove(conn, provider, key_id)
    conn.commit()


# ═════════════════════════════════════════════════════════════════════
# Storage dashboard + export/import (ADR §9.2, DoD #5/#6)
# ═════════════════════════════════════════════════════════════════════


def storage_usage(conn: DictConnection) -> StorageUsage:
    return get_storage_usage(conn)


def run_purge(
    conn: DictConnection,
    *,
    operation: str,
    older_than_days: int | None = None,
    older_than: date | None = None,
    require_exported: bool = True,
    confirm: bool = False,
) -> PurgeResult:
    """Dispatches to app.services.storage's three purge functions —
    consumed, never reimplemented (ARCH-DB owns the dry-run/confirm
    mechanics and is concurrently fixing a preview/delete race there).
    """
    if operation == "stale_bodies":
        return purge_stale_bodies(
            conn, older_than_days=older_than_days, require_exported=require_exported, confirm=confirm
        )
    if operation == "stale_unmatched":
        return purge_stale_unmatched_articles(conn, older_than_days=older_than_days, confirm=confirm)
    if operation == "matched":
        if older_than is None:
            raise ValidationError("older_than (an ISO date) is required for the 'matched' purge operation")
        return purge_matched_articles(conn, older_than=older_than, require_exported=require_exported, confirm=confirm)
    raise ValidationError(f"unknown purge operation {operation!r} — must be stale_bodies/stale_unmatched/matched")


EXPORT_DIR = "var/exports"


def run_export(
    conn: DictConnection,
    *,
    day_from: date | None = None,
    day_to: date | None = None,
    review_status: str | None = None,
    mark_exported: bool = True,
    path: str | None = None,
) -> ExportManifest:
    if path is None:
        ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
        path = f"{EXPORT_DIR}/export_{ts}.jsonl.gz"
    manifest = export_to_jsonl(
        conn, path, day_from=day_from, day_to=day_to, review_status=review_status, mark_exported=mark_exported
    )
    conn.commit()
    return manifest


def run_import(conn: DictConnection, *, path: str) -> ImportReport:
    report = import_from_jsonl(conn, path)
    conn.commit()
    return report


# ═════════════════════════════════════════════════════════════════════
# Source health + pipeline stage runs (frontend-contract.md §3.10)
# ═════════════════════════════════════════════════════════════════════

# The 5 pipeline stages, in the exact order/labels frontend-contract.md
# §3.10 names them. "bodies" (Get article text) has no existing
# orchestrator anywhere in this codebase — legacy itself runs
# `do_bodies=False` on every normal daily run (app.services.enrichment's
# own module docstring), and P1-INGEST's collector never fetches bodies
# either — so it is implemented here, minimally and boundedly, directly
# against the read-only app.sources plugin interface (the same one
# probe_source/diagnose_source above already use). The other four
# dispatch straight to app.services.enrichment.run_llm_enrichment,
# P3-LLM's already-built stage orchestrator — not reimplemented here.
PIPELINE_STAGES: tuple[tuple[str, str, str], ...] = (
    ("translate", "Translate", "Bengali headlines → English"),
    ("bodies", "Get article text", "Fetch the full article body for today's matched stories"),
    ("dedup", "Group duplicates", "Cluster near-duplicate coverage of the same event"),
    ("judge", "Score relevance", "LLM relevance scoring, 1-10"),
    ("topic", "Group topics", "Cluster scored articles into broad topics"),
)
_PIPELINE_STAGE_KEYS = {key for key, _, _ in PIPELINE_STAGES}


def _stage_status_setting_key(stage: str, day: date) -> str:
    return f"pipeline_stage_status:{stage}:{day.isoformat()}"


def _record_stage_status(
    conn: DictConnection, stage: str, *, status: str, detail: str, day: date | None = None
) -> None:
    day = day or date.today()
    set_setting(
        conn,
        _stage_status_setting_key(stage, day),
        json.dumps({"status": status, "detail": detail, "ran_at": datetime.now(timezone.utc).isoformat()}),
    )
    conn.commit()


def pipeline_stage_statuses(conn: DictConnection, day: date | None = None) -> list[dict[str, Any]]:
    day = day or date.today()
    out = []
    for key, name, desc in PIPELINE_STAGES:
        raw = get_setting(conn, _stage_status_setting_key(key, day), default="")
        if raw:
            try:
                data = json.loads(raw)
            except json.JSONDecodeError:
                data = {}
            out.append(
                {
                    "key": key,
                    "name": name,
                    "desc": desc,
                    "status": data.get("status", "done"),
                    "ran_at": data.get("ran_at"),
                    "detail": data.get("detail", ""),
                }
            )
        else:
            out.append({"key": key, "name": name, "desc": desc, "status": "not_run", "ran_at": None, "detail": ""})
    return out


BODIES_STAGE_MAX = 2000
BODIES_STAGE_WORKERS = 12


def _run_bodies_stage(
    conn: DictConnection, *, limit: int = BODIES_STAGE_MAX, progress: Any | None = None
) -> dict[str, int]:
    """Fetch full article text for today's matched articles.

    Two bugs fixed here, both observed on a live run (2026-08-14, 173
    matched articles reported "attempted 100"):

    1. The cap was a hardcoded `limit=100` with no caller able to raise
       it and nothing in the UI saying a limit had been applied — so 73
       of 173 matched articles silently never got a body, and the
       downstream topic stage (which reads `content_text`) quietly saw a
       truncated corpus. The ceiling is now `BODIES_STAGE_MAX`, high
       enough to cover a real day's matched set, and the returned detail
       reports `capped` explicitly whenever it does bind, so a truncated
       run can never again look identical to a complete one.

    2. Fetches ran strictly sequentially — one blocking HTTP round-trip
       per article against ~45 different hosts. These are independent by
       construction and are now issued through a bounded pool. The
       process-wide ceiling in `app.sources.http.GLOBAL_SEMAPHORE` still
       applies underneath, so total concurrent connections stay bounded
       exactly as before.

    Threading contract: only the network call runs on the pool. Every DB
    write happens on the calling thread after the pool drains, because
    `conn` is a single psycopg connection and is not thread-safe.
    """
    import threading as _threading
    from concurrent.futures import ThreadPoolExecutor

    from app.sources.base import FeedItem
    from app.sources.registry import all_sources

    _ensure_registry_loaded()
    by_display_name = {getattr(cls, "display_name", cls.key): cls for cls in all_sources().values()}

    rows = conn.execute(
        "SELECT id, url, title, source FROM articles "
        "WHERE day = CURRENT_DATE AND review_status = 'pending' AND body_status = 'not_attempted' "
        "ORDER BY id LIMIT %s",
        (limit + 1,),
    ).fetchall()
    capped = len(rows) > limit
    rows = rows[:limit]

    # One source instance per (thread, source) — a `requests.Session` is
    # not documented as thread-safe, so sharing one across workers would
    # be a real (if intermittent) hazard, while building a fresh one per
    # ARTICLE would throw away connection reuse entirely. Thread-local
    # instances give both safety and keep-alive.
    local = _threading.local()

    def _instance_for(source_cls: Any) -> Any:
        cache = getattr(local, "sources", None)
        if cache is None:
            cache = {}
            local.sources = cache
        key = getattr(source_cls, "key", source_cls.__name__)
        inst = cache.get(key)
        if inst is None:
            inst = source_cls()
            cache[key] = inst
        return inst

    def _fetch_one(row: dict[str, Any]) -> tuple[dict[str, Any], str | None, bool]:
        """Returns (row, body_text_or_None, skipped_no_plugin)."""
        source_cls = by_display_name.get(row["source"])
        if source_cls is None:
            return row, None, True
        try:
            instance = _instance_for(source_cls)
            item = FeedItem(url=row["url"], title=row["title"], source_key=getattr(source_cls, "key", ""))
            return row, instance.fetch_article(item).body_text, False
        except SourceError:
            return row, None, False
        except Exception:  # noqa: BLE001 - one bad page must never abort the stage
            return row, None, False

    results: list[tuple[dict[str, Any], str | None, bool]] = []
    if rows:
        done = 0
        done_lock = _threading.Lock()

        def _fetch_and_report(row: dict[str, Any]) -> tuple[dict[str, Any], str | None, bool]:
            nonlocal done
            out = _fetch_one(row)
            if progress is not None:
                with done_lock:
                    done += 1
                    seen = done
                try:
                    progress(seen, len(rows), "Fetching article text")
                except Exception:  # noqa: BLE001 - reporting must never fail the stage
                    pass
            return out

        with ThreadPoolExecutor(max_workers=min(BODIES_STAGE_WORKERS, len(rows))) as pool:
            results = list(pool.map(_fetch_and_report, rows))

    fetched = failed = skipped = 0
    for row, body_text, no_plugin in results:
        if no_plugin:
            skipped += 1
            continue
        if body_text is None:
            conn.execute("UPDATE articles SET body_status = 'failed' WHERE id = %s", (row["id"],))
            failed += 1
            continue
        conn.execute(
            "INSERT INTO article_bodies (article_id, content_text, fetched_at) VALUES (%s, %s, now()) "
            "ON CONFLICT (article_id) DO UPDATE SET content_text = EXCLUDED.content_text, fetched_at = now()",
            (row["id"], body_text),
        )
        conn.execute("UPDATE articles SET body_status = 'fetched' WHERE id = %s", (row["id"],))
        fetched += 1
    conn.commit()
    return {
        "attempted": len(rows),
        "fetched": fetched,
        "failed": failed,
        "skipped_no_plugin": skipped,
        "capped": 1 if capped else 0,
    }


def run_pipeline_stage(
    conn: DictConnection,
    stage_key: str,
    *,
    progress: Any | None = None,
) -> dict[str, str]:
    """Run ONE enrichment stage.

    `progress(done, total, step)` is forwarded to the enrichment stage
    so callers can render live x/n counts instead of an opaque
    "running…". Optional: every existing caller that omits it behaves
    exactly as before.
    """
    if stage_key not in _PIPELINE_STAGE_KEYS:
        raise ValidationError(f"unknown pipeline stage {stage_key!r} — must be one of {sorted(_PIPELINE_STAGE_KEYS)}")
    try:
        if stage_key == "bodies":
            result = _run_bodies_stage(conn, progress=progress)
            detail = (
                f"attempted {result['attempted']}, fetched {result['fetched']}, "
                f"failed {result['failed']}, skipped (no plugin) {result['skipped_no_plugin']}"
            )
            if result.get("capped"):
                # Never let a truncated run render identically to a
                # complete one — that is what hid the old limit=100.
                detail += f" — CAPPED at {BODIES_STAGE_MAX}, more remain"
        else:
            # Lazy import — see _build_client_for_feature's docstring for
            # why app.services.llm/.enrichment are never imported at
            # module level here: sys.modules caches the module either
            # way, so a test that does `from app.services import
            # enrichment; monkeypatch.setattr(enrichment, ...)` still
            # patches the exact object this resolves to.
            from app.services import enrichment

            enrichment_result = enrichment.run_llm_enrichment(conn, stages={stage_key}, progress=progress)
            detail = f"processed {enrichment_result.enriched} article(s)"
            if enrichment_result.skipped_stages:
                detail += f"; skipped (no provider configured): {', '.join(enrichment_result.skipped_stages)}"
    except Exception as exc:
        _record_stage_status(conn, stage_key, status="error", detail=str(exc))
        raise
    _record_stage_status(conn, stage_key, status="done", detail=detail)
    return {"stage": stage_key, "status": "done", "detail": detail}


def run_pipeline_all(conn: DictConnection) -> list[dict[str, str]]:
    """Runs all 5 stages in order. Each stage's failure is isolated —
    one broken stage is recorded and the rest still run — mirroring
    app.services.collector's per-source isolation philosophy at the
    pipeline-stage level instead of the per-source level.
    """
    results = []
    for stage_key, _, _ in PIPELINE_STAGES:
        try:
            results.append(run_pipeline_stage(conn, stage_key))
        except Exception as exc:  # noqa: BLE001 - deliberate per-stage isolation, see docstring
            results.append({"stage": stage_key, "status": "error", "detail": str(exc)})
    return results


# ═════════════════════════════════════════════════════════════════════
# Fetch (P1) trigger — "run collection now" (BUGFIX: "Run all stages"
# only ever ran the 5 LLM enrichment stages — nothing anywhere called
# app.services.collector.run_collection except app.batch.run_daily.main,
# itself only ever invoked by the daily GitHub Actions cron. See that
# module's own docstring: trigger="scheduled" is "the only real
# entrypoint... 'manual' [is] the schema's own default, meant for a
# future admin-triggered 'run now' button this module doesn't
# implement." This is that button's service layer.)
#
# Date selection is deliberately NOT exposed anywhere in this surface
# (owner decision, not re-litigated here): `run_collection` always
# fetches for "today" (`app.sources._util.today_dhaka()`, Dhaka-local) —
# only 17 of 45 sources could ever honor a specific past date; the rest
# are live current-only feeds, so a date picker would silently lie for
# the other 28.
#
# A real P1 fetch takes low minutes across ~45 sources
# (README-rebuild.md: "~2.5-minute full 45-source collection phase") —
# far too long to hold an HTTP request open for. Runs in a background
# `threading.Thread` instead: this is a single-worker free-tier deploy
# (mission constraint — no new dependency, no real job queue), so a
# plain daemon thread sharing the same process's memory is both the
# simplest mechanism available and sufficient. The in-process
# `_fetch_state` dict below is visible to every request handler in that
# same single worker without needing a pubsub/queue of its own — it
# would NOT be sufficient across multiple worker processes, which this
# deploy does not run (flagged in this task's report as the thing that
# would need to change first).
#
# `_fetch_state` is intentionally a single global slot, not a dict keyed
# by run id: only one manual fetch makes sense in flight at a time on a
# single-worker deploy (a second concurrent 45-source fetch would just
# contend with the first over the same `app.sources.http.
# GLOBAL_SEMAPHORE` and DB connection count for no benefit) —
# `start_fetch_run` returns the existing in-progress snapshot unchanged,
# rather than starting a second one stacked behind it, if one is already
# running.
#
# Real per-source progress while the fetch is running comes from
# `app.services.collector.run_collection`'s new `progress_cb` parameter
# (this task's one addition to that module's internals — see its and
# `fetch_all`'s docstrings for the full reasoning). The short version:
# `fetch_all`'s old behavior blocked until EVERY source's discovery had
# finished before returning anything, so even the existing
# `runs.heartbeat_at`/`runs.details` RT-1 infrastructure had nothing to
# report until the whole ~2.5-minute fan-out was already over —
# "instant done" all over again, just with a multi-minute wait first
# instead of a zero-second one. `progress_cb` fires per-source, as each
# one's fetch+write actually completes (interleaved with the other
# sources still in flight), in real time — the mechanism that makes
# `fetch_status()` below genuinely live rather than "pending" then
# instantly "done".

_fetch_lock = threading.Lock()
_fetch_state: dict[str, Any] = {"status": "idle"}


def _fetch_source_progress_entry(source_result: Any) -> dict[str, Any]:
    return {
        "name": source_result.display_name,
        "status": source_result.status,
        "fetched": source_result.fetched,
        "inserted": source_result.inserted,
        "duplicate": source_result.duplicate,
        "error": source_result.error,
    }


def _run_fetch_in_background(trigger: str) -> None:
    """Runs on the background thread `start_fetch_run` spawns — never on
    the request thread. Opens its OWN connection
    (`app.db.connection.get_connection`): the HTTP request that
    triggered this has already returned (that's the whole point) by the
    time most of this executes, so it cannot reuse that request's
    connection. Updates the module-level `_fetch_state` under
    `_fetch_lock` as `run_collection`'s `progress_cb` fires per source,
    then once more at the very end (or on failure).

    `app.services.collector._filter_enabled` (a "private", underscore
    module function) is called directly here rather than duplicating its
    one-query logic — this file already reads collector.py's public
    surface elsewhere; reusing this one existing helper for an accurate
    "N sources will actually run today" denominator (excluding
    admin-disabled sources up front) is simpler and less risky than a
    second, slightly-different COUNT query that could drift out of sync
    with it.
    """
    from app.db.connection import get_connection
    from app.services.collector import _filter_enabled, run_collection
    from app.sources.registry import all_sources, load_all

    def _on_progress(source_result: Any) -> None:
        with _fetch_lock:
            if _fetch_state.get("status") == "running":
                _fetch_state.setdefault("sources", []).append(_fetch_source_progress_entry(source_result))

    try:
        load_all()
        registry = all_sources()
        with get_connection() as conn:
            effective = _filter_enabled(conn, registry)
            with _fetch_lock:
                if _fetch_state.get("status") == "running":
                    _fetch_state["total_sources"] = len(effective)
            result = run_collection(conn, source_classes=effective, trigger=trigger, progress_cb=_on_progress)
        with _fetch_lock:
            _fetch_state.update(
                status="done",
                run_id=result.run_id,
                finished_at=datetime.now(timezone.utc).isoformat(),
                total_fetched=result.total_fetched,
                total_inserted=result.total_inserted,
                total_duplicate=result.total_duplicate,
            )
    except Exception as exc:  # noqa: BLE001 - background thread: must never raise into the void, see docstring
        with _fetch_lock:
            _fetch_state.update(
                status="error",
                finished_at=datetime.now(timezone.utc).isoformat(),
                error=str(exc),
            )


def start_fetch_run(*, trigger: str = "manual") -> dict[str, Any]:
    """`POST /health/pipeline/fetch/run`'s service function. Starts a
    real `app.services.collector.run_collection` for TODAY
    (Dhaka-local — see module-section docstring on why there is no date
    parameter) in a background thread and returns immediately with a
    "running" snapshot — never blocks the request for the full fetch
    duration. Returns the existing in-progress snapshot unchanged,
    without starting a second run, if one is already running.
    """
    with _fetch_lock:
        already_running = _fetch_state.get("status") == "running"
        if not already_running:
            _fetch_state.clear()
            _fetch_state.update(
                status="running",
                run_id=None,
                trigger=trigger,
                started_at=datetime.now(timezone.utc).isoformat(),
                finished_at=None,
                total_sources=0,
                sources=[],
                total_fetched=0,
                total_inserted=0,
                total_duplicate=0,
                error=None,
            )
    # `fetch_status()` acquires `_fetch_lock` itself and MUST be called
    # after the `with` block above has released it — `threading.Lock` is
    # not reentrant, so calling it from inside that block (a real bug
    # caught by this task's own tests) self-deadlocks the very next
    # "fetch already running" request.
    if not already_running:
        thread = threading.Thread(target=_run_fetch_in_background, args=(trigger,), daemon=True)
        thread.start()
    return fetch_status()


def fetch_status() -> dict[str, Any]:
    """`GET /health/pipeline/fetch/status`'s service function — a live
    read of the in-process `_fetch_state`, never a DB query: the whole
    point is to surface progress the `runs` row itself cannot show until
    `_finalize_run` writes `details` at the very end (see
    `app.services.collector.run_collection`'s docstring). Safe to call
    at any time, including before any fetch has ever run
    (`status="idle"`, empty `sources`).
    """
    with _fetch_lock:
        state = dict(_fetch_state)
        sources = list(state.get("sources", []))
    state["sources"] = sources
    state["completed"] = len(sources)
    state.setdefault("status", "idle")
    state.setdefault("total_sources", 0)
    state.setdefault("trigger", None)
    state.setdefault("run_id", None)
    state.setdefault("started_at", None)
    state.setdefault("finished_at", None)
    state.setdefault("total_fetched", 0)
    state.setdefault("total_inserted", 0)
    state.setdefault("total_duplicate", 0)
    state.setdefault("error", None)
    return state


# ═════════════════════════════════════════════════════════════════════
# Match (P2) trigger — "keyword-match today's fetched articles now"
# (BUGFIX mission, phase selection requirement). P2 is genuinely a
# separate step from P1 in this codebase: `run_collection`
# (app.services.collector) never calls into app.services.matching at
# all — its own module docstring is explicit that it "deliberately does
# NOT compute sections/keywords/review_status... that is
# app.services.matching's job". `rematch_today` (above, in the Keywords
# section of this file) already IS that job, already scoped to today's
# `pending`/`filtered_out` rows, already safe to re-run repeatedly (never
# touches an analyst's approve/reject decision) — this is a thin,
# same-name wrapper purely so the pipeline/phase surface below has one
# consistently-named function per phase to call, not a reimplementation.
# ═════════════════════════════════════════════════════════════════════


def run_match_phase(conn: DictConnection) -> dict[str, int]:
    return rematch_today(conn)


def health_sources(conn: DictConnection) -> list[dict[str, Any]]:
    """`healthRows[]` (§3.10) — one row per configured source, joined
    against the most recent collection run's per-source `details` JSONB
    (app.services.collector's `runs.details`, keyed by display name).
    """
    sources_rows = list_sources(conn)
    latest_run = conn.execute("SELECT details FROM runs ORDER BY started_at DESC LIMIT 1").fetchone()
    details: dict[str, Any] = dict(latest_run["details"]) if latest_run and latest_run["details"] else {}

    rows = []
    for source in sources_rows:
        d = details.get(source["name"], {})
        status_raw = d.get("status")
        fetched = int(d.get("fetched", 0))
        inserted = int(d.get("inserted", 0))
        if not source["enabled"]:
            status, note = "disabled", "source disabled"
        elif status_raw == "failed":
            status, note = "error", f"error: {d.get('error', '')}"
        elif status_raw == "ok" and fetched > 0 and inserted == 0:
            status, note = "quiet", f"fetched {fetched} · {fetched} already stored"
        elif status_raw == "ok" and fetched == 0:
            status, note = "quiet", "0 fetched this run"
        elif status_raw == "ok":
            status, note = "healthy", ""
        elif status_raw == "empty":
            status, note = "quiet", "no articles found this run"
        else:
            status, note = "quiet", "no run data yet"
        rows.append(
            {
                "id": source["id"],
                "name": source["name"],
                "method": source["methods"][0] if source["methods"] else "",
                "fetched": fetched,
                "kept": inserted,
                "status": status,
                "note": note,
            }
        )
    return rows


def health_overview(conn: DictConnection) -> dict[str, Any]:
    """The KPI tiles §3.10 flags as unbound literals in the mock —
    computed for real from `health_sources`/`health_runs`, never a fixed
    display value.
    """
    rows = health_sources(conn)
    latest_run = conn.execute("SELECT started_at, total_matched FROM runs ORDER BY started_at DESC LIMIT 1").fetchone()
    return {
        "healthy": sum(1 for r in rows if r["status"] == "healthy"),
        "needs_attention": sum(1 for r in rows if r["status"] == "error"),
        "fetched_last_run": sum(r["fetched"] for r in rows),
        "matched_last_run": latest_run["total_matched"] if latest_run else 0,
        "last_run_at": latest_run["started_at"].isoformat() if latest_run else None,
        "stages": pipeline_stage_statuses(conn),
    }


# RT-1-residual (docs/red-team-findings.md,
# tests/test_redteam_orphaned_runs.py): app.db.reaper.reap_stale_runs
# (DEPLOY, 59dbae7) self-heals a hard-killed run's abandoned
# status='running' row on the NEXT batch invocation, but between the
# kill and that next run, nothing surfaced heartbeat_at at all -- a
# stuck row read exactly like a healthy one. Mirrors
# app.db.reaper.DEFAULT_STALE_AFTER so "stale" means the same thing here
# as it does to the thing that eventually reaps it.
_HEARTBEAT_STALE_AFTER = timedelta(minutes=30)


def health_runs(conn: DictConnection, limit: int = 5) -> list[dict[str, Any]]:
    rows = conn.execute(
        "SELECT id, started_at, heartbeat_at, trigger, total_new, total_matched, status "
        "FROM runs ORDER BY started_at DESC LIMIT %s",
        (limit,),
    ).fetchall()
    now = datetime.now(timezone.utc)
    out = []
    for r in rows:
        heartbeat_at = r["heartbeat_at"]
        last_beat = heartbeat_at or r["started_at"]
        heartbeat_stale = bool(
            r["status"] == "running" and last_beat is not None and (now - last_beat) > _HEARTBEAT_STALE_AFTER
        )
        out.append(
            {
                "id": r["id"],
                "when": r["started_at"].isoformat() if r["started_at"] else "",
                "trigger": r["trigger"],
                "new_count": r["total_new"],
                "matched_count": r["total_matched"],
                "status": r["status"],
                "heartbeat_at": heartbeat_at.isoformat() if heartbeat_at else None,
                "heartbeat_stale": heartbeat_stale,
            }
        )
    return out


# ═════════════════════════════════════════════════════════════════════
# Consolidated run (BUGFIX-NAV mission): "Run collection" used to be a
# CLIENT-SIDE conductor — frontend/js/screens/dashboard.js's
# runCollection() called fetch, then (if the closure's `alive` flag was
# still true) match, then (same check) enrich, sequentially, awaiting
# each phase's own polling loop on the browser's call stack. The moment
# an admin navigated away mid-fetch, that screen's destroy() flipped
# `alive = false`; runFetchPhase()'s poll loop hit `if (!alive) return`
# and bailed BEFORE fetch's status ever went terminal, so runCollection()
# resumed with `alive` false and `if (alive && pipeline.match)` never
# ran. Fetch itself kept running fine server-side (it was already a
# background thread — see "Fetch (P1) trigger" above) and completed with
# real inserted rows; match and enrich just never got triggered, because
# the ONLY thing that decided "call match next" was JS holding a
# reference to a screen that had already been torn down. Real-world
# result: thousands of genuinely-fetched articles, almost none matched.
#
# Fix: the phase-sequencing decision moves entirely into THIS background
# thread — the identical `threading.Thread` + module-level dict-and-lock
# pattern `_fetch_state`/`start_fetch_run`/`fetch_status` above already
# established (same single-worker-deploy justification, unchanged). Once
# `start_run` returns, fetch -> match -> enrich proceeds to completion
# with zero further involvement from any HTTP request or browser tab.
# The frontend's job shrinks to "start it, then poll to render" — see
# dashboard.js's rewritten runCollection().
#
# Each phase's failure is isolated from the next (mirrors
# run_pipeline_all's per-stage isolation, and _run_fetch_in_background's
# own per-run isolation): a failed fetch does not prevent match or
# enrich from being attempted, matching the previous UI-level contract
# ("one phase failing does not block the next").
#
# `_run_state` is, like `_fetch_state`, a single global slot rather than
# a dict keyed by run id — only one manual collection run makes sense in
# flight at a time on this single-worker deploy. `start_run` returns the
# existing in-progress snapshot unchanged if one is already running,
# rather than stacking a second one behind it.

_run_lock = threading.Lock()
_run_state: dict[str, Any] = {"status": "idle", "phase": "idle"}


def _new_fetch_phase_state() -> dict[str, Any]:
    return {
        "status": "pending",
        "detail": "",
        "sources": [],
        "total_sources": 0,
        "completed": 0,
        "total_fetched": 0,
        "total_inserted": 0,
        "total_duplicate": 0,
        "run_id": None,
        "error": None,
    }


def _new_match_phase_state() -> dict[str, Any]:
    return {"status": "pending", "detail": "", "matched": 0, "filtered_out": 0, "total": 0, "error": None}


def _now_iso() -> str:
    return datetime.now(timezone.utc).isoformat()


def _new_enrich_phase_state() -> dict[str, Any]:
    """Per-stage state, including LIVE progress fields.

    `done`/`total`/`step` are what turn an opaque "running…" into
    "Scoring relevance · 430/1376". They are updated from the stage's
    own progress callback while it runs, so an operator can tell a slow
    stage apart from a stuck one — previously indistinguishable, which
    is exactly how a genuinely hung judge stage went unnoticed.
    """
    return {
        "status": "pending",
        "stages": [
            {
                "key": key,
                "name": name,
                "desc": desc,
                "status": "pending",
                "detail": "",
                "done": 0,
                "total": 0,
                "step": "",
                "started_at": None,
                "finished_at": None,
            }
            for key, name, desc in PIPELINE_STAGES
        ],
    }


def _run_fetch_subphase(conn: DictConnection, *, trigger: str) -> None:
    from app.services.collector import _filter_enabled, run_collection
    from app.sources.registry import all_sources, load_all

    def _on_progress(source_result: Any) -> None:
        with _run_lock:
            if _run_state.get("status") == "running" and _run_state.get("fetch") is not None:
                _run_state["fetch"]["sources"].append(_fetch_source_progress_entry(source_result))
                _run_state["fetch"]["completed"] = len(_run_state["fetch"]["sources"])

    with _run_lock:
        _run_state["phase"] = "fetch"
        _run_state["fetch"]["status"] = "running"
    try:
        load_all()
        registry = all_sources()
        effective = _filter_enabled(conn, registry)
        with _run_lock:
            _run_state["fetch"]["total_sources"] = len(effective)
        result = run_collection(conn, source_classes=effective, trigger=trigger, progress_cb=_on_progress)
        with _run_lock:
            _run_state["fetch"].update(
                status="done",
                run_id=result.run_id,
                total_fetched=result.total_fetched,
                total_inserted=result.total_inserted,
                total_duplicate=result.total_duplicate,
                detail=f"{result.total_inserted} new article(s) across {len(effective)} source(s)",
            )
    except Exception as exc:  # noqa: BLE001 - per-phase isolation, match/enrich must still get a chance to run
        with _run_lock:
            _run_state["fetch"].update(status="error", error=str(exc), detail=str(exc))


def _run_match_subphase(conn: DictConnection) -> None:
    with _run_lock:
        _run_state["phase"] = "match"
        _run_state["match"]["status"] = "running"
    try:
        result = run_match_phase(conn)
        with _run_lock:
            _run_state["match"].update(
                status="done",
                matched=result["matched"],
                filtered_out=result["filtered_out"],
                total=result["total"],
                detail=f"{result['matched']} matched, {result['filtered_out']} filtered out (of {result['total']})",
            )
    except Exception as exc:  # noqa: BLE001 - per-phase isolation, enrich must still get a chance to run
        with _run_lock:
            _run_state["match"].update(status="error", error=str(exc), detail=str(exc))


# The enrichment DAG, as WAVES of stages that may run concurrently.
#
# The five stages were run strictly one after another, which was safe
# but slower than it needs to be — several of them have no relationship
# to each other at all. The real dependencies are on COLUMNS:
#
#   translate  reads title            -> writes title_en
#   bodies     reads url              -> writes article_bodies (no LLM)
#   dedup      reads title_en         -> writes dup_group/is_primary/also_in
#   judge      reads title_en         -> writes llm_score/llm_reason
#   topic      reads llm_score >= 5   -> writes topic_tag/topic_parent_id
#
# So `bodies` shares nothing with `translate` (different source column,
# different destination table, and it is pure HTTP with no LLM budget to
# contend for). `dedup` and `judge` both read title_en but write
# DISJOINT columns, so they cannot race each other either. `topic` must
# follow `judge` because it filters on the scores judge writes.
#
# Stages inside a wave run concurrently; waves run in order. Anything
# not listed here still runs, appended as its own final wave, so adding
# a stage to PIPELINE_STAGES can never silently drop it.
ENRICH_WAVES: tuple[tuple[str, ...], ...] = (
    ("translate", "bodies"),
    ("dedup", "judge"),
    ("topic",),
)


def _stage_entry(stage_key: str) -> dict[str, Any] | None:
    """Caller must hold `_run_lock`."""
    enrich = _run_state.get("enrich")
    if not enrich:
        return None
    return next((s for s in enrich.get("stages", []) if s["key"] == stage_key), None)


def _run_one_enrich_stage(stage_key: str) -> bool:
    """Run one stage on its OWN connection. Returns True on success.

    Its own connection, not the caller's: stages in a wave run on
    separate threads and a psycopg connection is not thread-safe.
    Sharing one would corrupt the protocol under concurrency — the
    failure mode is silent and awful, so each thread gets its own.

    Never raises. A stage that fails is recorded and the rest of the run
    continues, matching the per-source isolation the collector already
    uses one level down.
    """
    from app.db.connection import get_connection

    def _report(done: int, total: int, step: str) -> None:
        with _run_lock:
            entry = _stage_entry(stage_key)
            if entry is not None:
                entry["done"] = int(done)
                entry["total"] = int(total)
                entry["step"] = str(step)

    with _run_lock:
        entry = _stage_entry(stage_key)
        if entry is not None:
            entry.update(status="running", done=0, total=0, step="", started_at=_now_iso())
    try:
        with get_connection() as stage_conn:
            result = run_pipeline_stage(stage_conn, stage_key, progress=_report)
        with _run_lock:
            entry = _stage_entry(stage_key)
            if entry is not None:
                entry.update(status=result["status"], detail=result["detail"], finished_at=_now_iso())
        return True
    except Exception as exc:  # noqa: BLE001 - per-stage isolation
        logger.warning("enrich_stage_failed", extra={"stage": stage_key, "error": str(exc)[:300]})
        with _run_lock:
            entry = _stage_entry(stage_key)
            if entry is not None:
                entry.update(status="error", detail=str(exc)[:300], finished_at=_now_iso())
        return False


def _run_enrich_subphase(conn: DictConnection) -> None:
    """Run the enrichment DAG. `conn` is intentionally unused — each
    stage opens its own (see `_run_one_enrich_stage`) — but is kept in
    the signature so the phase functions stay uniform.
    """
    del conn  # each stage owns its connection; see _run_one_enrich_stage
    from concurrent.futures import ThreadPoolExecutor

    with _run_lock:
        _run_state["phase"] = "enrich"

    configured = [key for key, _, _ in PIPELINE_STAGES]
    planned = [tuple(s for s in wave if s in configured) for wave in ENRICH_WAVES]
    # Anything in PIPELINE_STAGES that no wave claims still has to run.
    claimed = {s for wave in planned for s in wave}
    leftovers = tuple(s for s in configured if s not in claimed)
    if leftovers:
        logger.info("enrich_unclaimed_stages_appended", extra={"stages": list(leftovers)})
        planned.append(leftovers)

    any_error = False
    for wave in planned:
        if not wave:
            continue
        if len(wave) == 1:
            any_error |= not _run_one_enrich_stage(wave[0])
            continue
        with ThreadPoolExecutor(max_workers=len(wave)) as pool:
            for ok in pool.map(_run_one_enrich_stage, wave):
                any_error |= not ok

    with _run_lock:
        _run_state["enrich"]["status"] = "error" if any_error else "done"


def _run_pipeline_in_background(*, fetch: bool, match: bool, enrich: bool, trigger: str) -> None:
    """Runs on the background thread `start_run` spawns — see this
    module section's docstring for the bug this closes. Opens its own
    connection, same reasoning as `_run_fetch_in_background`: the HTTP
    request that triggered this has already returned by the time most of
    this executes.
    """
    from app.db.connection import get_connection

    try:
        with get_connection() as conn:
            if fetch:
                _run_fetch_subphase(conn, trigger=trigger)
            if match:
                _run_match_subphase(conn)
            if enrich:
                _run_enrich_subphase(conn)
        with _run_lock:
            _run_state["status"] = "done"
            _run_state["phase"] = "done"
            _run_state["finished_at"] = datetime.now(timezone.utc).isoformat()
    except Exception as exc:  # noqa: BLE001 - background thread: must never raise into the void
        with _run_lock:
            _run_state["status"] = "error"
            _run_state["phase"] = "error"
            _run_state["finished_at"] = datetime.now(timezone.utc).isoformat()
            _run_state["error"] = str(exc)


def start_run(
    *, fetch: bool = True, match: bool = True, enrich: bool = True, trigger: str = "manual"
) -> dict[str, Any]:
    """`POST /health/pipeline/run`'s service function — the single entry
    point for "Run collection". Starts fetch/match/enrich (whichever are
    requested) in ONE background thread, in that order, and returns
    immediately with a "running" snapshot. Once this returns, nothing
    about the rest of the run depends on the calling request, or any
    browser, still existing — see this module section's docstring.
    Returns the existing in-progress snapshot unchanged, without starting
    a second run, if one is already running.
    """
    if not fetch and not match and not enrich:
        raise ValidationError("at least one phase (fetch/match/enrich) must be selected")
    with _run_lock:
        already_running = _run_state.get("status") == "running"
        if not already_running:
            _run_state.clear()
            _run_state.update(
                status="running",
                phase="idle",
                trigger=trigger,
                requested={"fetch": fetch, "match": match, "enrich": enrich},
                started_at=datetime.now(timezone.utc).isoformat(),
                finished_at=None,
                fetch=_new_fetch_phase_state() if fetch else None,
                match=_new_match_phase_state() if match else None,
                enrich=_new_enrich_phase_state() if enrich else None,
                error=None,
            )
    # `run_status()` acquires `_run_lock` itself — must be called after
    # the `with` block above releases it, same reentrancy reasoning as
    # `start_fetch_run`'s own docstring.
    if not already_running:
        thread = threading.Thread(
            target=_run_pipeline_in_background,
            kwargs={"fetch": fetch, "match": match, "enrich": enrich, "trigger": trigger},
            daemon=True,
        )
        thread.start()
    return run_status()


def run_status() -> dict[str, Any]:
    """`GET /health/pipeline/run/status`'s service function — a live read
    of the in-process `_run_state`, the one status surface the frontend
    polls to RENDER the whole fetch->match->enrich run. Purely
    observational: nothing this function (or its caller) does changes
    what the background thread does next — see this module section's
    docstring. Safe to call at any time, including before any run has
    ever started (`status="idle"`).
    """
    with _run_lock:
        state = copy.deepcopy(_run_state)
    state.setdefault("status", "idle")
    state.setdefault("phase", "idle")
    state.setdefault("trigger", None)
    state.setdefault("requested", None)
    state.setdefault("started_at", None)
    state.setdefault("finished_at", None)
    state.setdefault("fetch", None)
    state.setdefault("match", None)
    state.setdefault("enrich", None)
    state.setdefault("error", None)
    return state
